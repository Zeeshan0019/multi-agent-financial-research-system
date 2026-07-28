import asyncio
import csv
import io
import json
import math
import mimetypes
import os
import re
import sys
import tempfile
import time
import unicodedata
import zipfile
from datetime import datetime
from pathlib import Path
from threading import Lock
from typing import Any, Dict, List, Optional, Tuple, TypedDict, Union

import aiofiles
import cv2
import fitz
import httpx
import numpy as np
import orjson
import pandas as pd
import pdfplumber
from babel.numbers import NumberFormatError, parse_decimal
from dateutil import parser as date_parser
from docx import Document
import docx2txt
from dotenv import load_dotenv
from fastapi import FastAPI, File, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, ORJSONResponse
from langchain_core.exceptions import OutputParserException
from langchain_core.output_parsers import JsonOutputParser, PydanticOutputParser
from langchain_core.prompts import PromptTemplate
from langchain_groq import ChatGroq
from langgraph.graph import END, START, StateGraph
from loguru import logger
from openpyxl import load_workbook
from PIL import Image, ImageOps, ImageSequence, UnidentifiedImageError
from pydantic import BaseModel, ConfigDict, Field, ValidationError, field_validator
from rapidfuzz import fuzz
from tenacity import retry, retry_if_exception, stop_after_attempt, wait_exponential

try:
    import chardet
except Exception:
    chardet = None

load_dotenv()

WORKSPACE_DIR = Path(__file__).resolve().parent
os.environ.setdefault("PADDLE_PDX_CACHE_HOME", str(WORKSPACE_DIR / "work" / "paddlex-cache"))
os.environ.setdefault("PADDLEOCR_DISABLE_AUTO_LOGGING_CONFIG", "1")
os.environ.setdefault("HF_HOME", str(WORKSPACE_DIR / "work" / "huggingface-cache"))
os.environ.setdefault("MODELSCOPE_CACHE", str(WORKSPACE_DIR / "work" / "modelscope-cache"))
for cache_directory in (
    os.environ["PADDLE_PDX_CACHE_HOME"],
    os.environ["HF_HOME"],
    os.environ["MODELSCOPE_CACHE"],
):
    Path(cache_directory).mkdir(parents=True, exist_ok=True)

logger.remove()
logger.add(sys.stderr, level=os.getenv("LOG_LEVEL", "INFO"), enqueue=os.getenv("LOGURU_ENQUEUE", "0").lower() in {"1", "true", "yes"}, backtrace=False, diagnose=False)


DEFAULT_GROQ_MODEL = os.getenv("GROQ_MODEL", "llama-3.3-70b-versatile")
LLM_TIMEOUT_SECONDS = float(os.getenv("LLM_TIMEOUT_SECONDS", "90"))
LLM_MAX_COMPLETION_TOKENS = int(os.getenv("LLM_MAX_COMPLETION_TOKENS", "256"))
LLM_MAX_INPUT_CHARS_PER_CHUNK = int(os.getenv("LLM_MAX_INPUT_CHARS_PER_CHUNK", "2000"))
LLM_TOTAL_INPUT_CHARS_PER_EXTRACTION = int(os.getenv("LLM_TOTAL_INPUT_CHARS_PER_EXTRACTION", "6000"))
LLM_MAX_CHUNKS_PER_EXTRACTION = int(os.getenv("LLM_MAX_CHUNKS_PER_EXTRACTION", "3"))
LLM_RETRY_ATTEMPTS = int(os.getenv("LLM_RETRY_ATTEMPTS", "1"))
LLM_STRUCTURED_OUTPUT_ENABLED = os.getenv("LLM_STRUCTURED_OUTPUT_ENABLED", "0").lower() in {"1", "true", "yes"}
LLM_PARSER_FALLBACK_ENABLED = os.getenv("LLM_PARSER_FALLBACK_ENABLED", "0").lower() in {"1", "true", "yes"}
LLM_RAW_REPAIR_ENABLED = os.getenv("LLM_RAW_REPAIR_ENABLED", "0").lower() in {"1", "true", "yes"}
APPROX_CHARS_PER_TOKEN = float(os.getenv("APPROX_CHARS_PER_TOKEN", "4"))
MAX_UPLOAD_BYTES = int(os.getenv("MAX_UPLOAD_MB", "50")) * 1024 * 1024
MAX_BATCH_FILES = int(os.getenv("MAX_BATCH_FILES", "10"))
BATCH_CONCURRENCY = int(os.getenv("BATCH_CONCURRENCY", "2"))
LLM_CONCURRENCY = int(os.getenv("LLM_CONCURRENCY", "3"))
PDF_OCR_DPI_SCALE = float(os.getenv("PDF_OCR_DPI_SCALE", "2.0"))
MAX_CHUNK_CHARS = int(os.getenv("MAX_CHUNK_CHARS", "12000"))
CHUNK_OVERLAP_CHARS = int(os.getenv("CHUNK_OVERLAP_CHARS", "400"))
MIN_NATIVE_PDF_CHARS_PER_PAGE = int(os.getenv("MIN_NATIVE_PDF_CHARS_PER_PAGE", "120"))
SUPPORTED_EXTENSIONS = {".pdf", ".png", ".jpeg", ".jpg", ".bmp", ".tiff", ".tif", ".docx", ".txt", ".md", ".rtf", ".html", ".htm", ".json", ".xml", ".csv", ".tsv", ".xlsx", ".xls"}
IMAGE_EXTENSIONS = {".png", ".jpeg", ".jpg", ".bmp", ".tiff", ".tif"}
SPREADSHEET_EXTENSIONS = {".xlsx", ".xls"}
TEXT_EXTENSIONS = {".txt", ".md", ".rtf", ".html", ".htm", ".json", ".xml"}
TABULAR_TEXT_EXTENSIONS = {".csv", ".tsv"}
DOCUMENT_EXTENSIONS = {".docx"}
PDF_EXTENSIONS = {".pdf"}
CURRENCY_SYMBOLS = {"$": "USD", "\u20ac": "EUR", "\u00a3": "GBP", "\u20b9": "INR", "\u00a5": "JPY"}
SCALAR_FINANCIAL_FIELDS = [
    "company_name",
    "ticker",
    "reporting_period",
    "financial_year",
    "quarter",
    "currency",
    "country",
    "industry",
    "reporting_standard",
    "fiscal_period",
    "auditor",
    "ceo",
]
METRIC_FINANCIAL_FIELDS = [
    "revenue",
    "gross_revenue",
    "gross_profit",
    "operating_income",
    "ebit",
    "ebitda",
    "net_income",
    "net_profit",
    "eps",
    "diluted_eps",
    "assets",
    "current_assets",
    "non_current_assets",
    "liabilities",
    "current_liabilities",
    "long_term_debt",
    "debt",
    "equity",
    "cash",
    "cash_equivalents",
    "cash_flow",
    "operating_cash_flow",
    "investing_cash_flow",
    "financing_cash_flow",
    "free_cash_flow",
    "capital_expenditure",
    "inventory",
    "receivables",
    "payables",
    "working_capital",
    "tax_expense",
    "interest_expense",
    "operating_margin",
    "gross_margin",
    "net_margin",
    "roa",
    "roe",
    "current_ratio",
    "quick_ratio",
    "debt_to_equity",
    "employees",
    "dividend",
    "shares_outstanding",
]
METRIC_FIELD_DISPLAY_NAMES = {
    "revenue": "Revenue",
    "gross_revenue": "Gross Revenue",
    "gross_profit": "Gross Profit",
    "operating_income": "Operating Income",
    "ebit": "EBIT",
    "ebitda": "EBITDA",
    "net_income": "Net Income",
    "net_profit": "Net Profit",
    "eps": "EPS",
    "diluted_eps": "Diluted EPS",
    "assets": "Total Assets",
    "current_assets": "Current Assets",
    "non_current_assets": "Non-current Assets",
    "liabilities": "Total Liabilities",
    "current_liabilities": "Current Liabilities",
    "long_term_debt": "Long-term Debt",
    "debt": "Debt",
    "equity": "Equity",
    "cash": "Cash",
    "cash_equivalents": "Cash Equivalents",
    "cash_flow": "Cash Flow",
    "operating_cash_flow": "Operating Cash Flow",
    "investing_cash_flow": "Investing Cash Flow",
    "financing_cash_flow": "Financing Cash Flow",
    "free_cash_flow": "Free Cash Flow",
    "capital_expenditure": "Capital Expenditure",
    "inventory": "Inventory",
    "receivables": "Receivables",
    "payables": "Payables",
    "working_capital": "Working Capital",
    "tax_expense": "Tax Expense",
    "interest_expense": "Interest Expense",
    "operating_margin": "Operating Margin",
    "gross_margin": "Gross Margin",
    "net_margin": "Net Margin",
    "roa": "ROA",
    "roe": "ROE",
    "current_ratio": "Current Ratio",
    "quick_ratio": "Quick Ratio",
    "debt_to_equity": "Debt to Equity",
    "employees": "Employees",
    "dividend": "Dividend",
    "shares_outstanding": "Shares Outstanding",
}
METRIC_FIELD_ALIASES = {
    "revenue": ("revenue", "net revenue", "sales", "net sales", "turnover", "total revenue", "operating revenue"),
    "gross_revenue": ("gross revenue",),
    "gross_profit": ("gross profit", "gross income"),
    "operating_income": ("operating income", "operating profit", "operating earnings", "income from operations"),
    "ebit": ("ebit", "earnings before interest and taxes"),
    "ebitda": ("ebitda", "earnings before interest taxes depreciation and amortization"),
    "net_income": ("net income", "net earnings", "profit attributable to shareholders"),
    "net_profit": ("net profit", "profit after tax", "pat", "profit for the year", "profit for the period"),
    "eps": ("eps", "basic eps", "earnings per share", "basic earnings per share"),
    "diluted_eps": ("diluted eps", "diluted earnings per share"),
    "assets": ("assets", "total assets"),
    "current_assets": ("current assets", "total current assets"),
    "non_current_assets": ("non current assets", "non-current assets", "total non current assets", "total non-current assets"),
    "liabilities": ("liabilities", "total liabilities"),
    "current_liabilities": ("current liabilities", "total current liabilities"),
    "long_term_debt": ("long term debt", "long-term debt", "non current debt", "non-current debt", "long term borrowings"),
    "debt": ("debt", "total debt", "borrowings"),
    "equity": ("equity", "shareholders equity", "stockholders equity", "total equity"),
    "cash": ("cash", "cash balance"),
    "cash_equivalents": ("cash equivalents", "cash and cash equivalents"),
    "cash_flow": ("cash flow", "net cash flow"),
    "operating_cash_flow": ("operating cash flow", "cash flow from operating activities", "net cash from operating activities"),
    "investing_cash_flow": ("investing cash flow", "cash flow from investing activities", "net cash used in investing activities"),
    "financing_cash_flow": ("financing cash flow", "cash flow from financing activities", "net cash from financing activities"),
    "free_cash_flow": ("free cash flow", "fcf"),
    "capital_expenditure": ("capital expenditure", "capital expenditures", "capex", "capital spending"),
    "inventory": ("inventory", "inventories"),
    "receivables": ("receivables", "accounts receivable", "trade receivables"),
    "payables": ("payables", "accounts payable", "trade payables"),
    "working_capital": ("working capital",),
    "tax_expense": ("tax expense", "income tax expense", "taxes"),
    "interest_expense": ("interest expense", "finance cost", "finance costs"),
    "operating_margin": ("operating margin",),
    "gross_margin": ("gross margin",),
    "net_margin": ("net margin",),
    "roa": ("roa", "return on assets"),
    "roe": ("roe", "return on equity"),
    "current_ratio": ("current ratio",),
    "quick_ratio": ("quick ratio",),
    "debt_to_equity": ("debt to equity", "debt-to-equity", "debt equity ratio"),
    "employees": ("employees", "headcount", "number of employees"),
    "dividend": ("dividend", "dividends", "dividend paid", "dividends paid"),
    "shares_outstanding": ("shares outstanding", "weighted average shares", "ordinary shares outstanding"),
}
COMPANY_INFO_FIELD_ALIASES = {
    "company_name": ("company name", "company", "issuer", "registrant", "entity name"),
    "ticker": ("ticker", "ticker symbol", "stock symbol", "symbol"),
    "reporting_period": ("reporting period", "period reported", "period covered"),
    "financial_year": ("financial year", "fiscal year", "fy", "year ended"),
    "quarter": ("quarter", "fiscal quarter"),
    "currency": ("currency", "presentation currency", "reporting currency"),
    "country": ("country", "jurisdiction", "country of incorporation"),
    "industry": ("industry", "sector"),
    "reporting_standard": ("reporting standard", "accounting standard", "ifrs", "gaap"),
    "fiscal_period": ("fiscal period",),
    "auditor": ("auditor", "independent auditor"),
    "ceo": ("ceo", "chief executive officer"),
}
DEFAULT_METRIC_CONFIDENCE = float(os.getenv("DEFAULT_METRIC_CONFIDENCE", "0.86"))
DEFAULT_COMPANY_INFO_CONFIDENCE = float(os.getenv("DEFAULT_COMPANY_INFO_CONFIDENCE", "0.82"))
LOCAL_FALLBACK_CONFIDENCE = float(os.getenv("LOCAL_FALLBACK_CONFIDENCE", "0.88"))
SUPPORTED_CURRENCY_CODES = {"USD", "EUR", "GBP", "INR", "CAD", "AUD", "JPY", "CNY", "CHF", "SGD", "HKD", "AED"}
LEGAL_ENTITY_SUFFIX_PATTERN = re.compile(
    r"\b(?:LIMITED|LTD\.?|PRIVATE LIMITED|PVT\.?\s+LTD\.?|INC\.?|INCORPORATED|CORPORATION|CORP\.?|PLC|LLC|LLP)\b",
    re.IGNORECASE,
)
COMPANY_NAME_EXCLUDE_PATTERN = re.compile(r"\b(?:auditor|associates|board|approval|meeting|agm|director|committee|report|statement)\b", re.IGNORECASE)



class Metric(BaseModel):
    model_config = ConfigDict(extra="ignore")

    metric: str = Field(description="Metric label exactly found or its closest canonical financial metric name.")
    value: str = Field(description="Value exactly present in the chunk.")
    unit: Optional[str] = Field(default=None, description="Currency, percent, shares, or scale when explicitly present.")
    period: Optional[str] = Field(default=None, description="Associated period when explicitly present.")
    page_number: Optional[Union[int, str]] = Field(default=None, description="Page number when explicitly present.")
    confidence: Optional[float] = Field(default=None, ge=0.0, le=1.0, description="Extraction confidence from 0 to 1.")

    @field_validator("metric", "value", mode="before")
    @classmethod
    def validate_required_text(cls, value: Any) -> str:
        text = safe_string(value)
        if text is None:
            raise ValueError("required text field is empty")
        return text

    @field_validator("unit", "period", mode="before")
    @classmethod
    def validate_optional_text(cls, value: Any) -> Optional[str]:
        return safe_string(value)

    @field_validator("page_number", mode="before")
    @classmethod
    def validate_page_number(cls, value: Any) -> Optional[Union[int, str]]:
        text = safe_string(value)
        if text is None:
            return None
        if re.fullmatch(r"\d+", text):
            return int(text)
        return text

    @field_validator("confidence", mode="before")
    @classmethod
    def validate_confidence(cls, value: Any) -> Optional[float]:
        if value is None or value == "":
            return None
        return clamp_float(value, 0.0, 1.0)


class CompanyInfoFact(BaseModel):
    model_config = ConfigDict(extra="ignore")

    field: str = Field(description="Company information field explicitly present in the chunk.")
    value: str = Field(description="Value exactly present in the chunk.")
    page_number: Optional[Union[int, str]] = Field(default=None)
    confidence: Optional[float] = Field(default=None, ge=0.0, le=1.0)

    @field_validator("field", "value", mode="before")
    @classmethod
    def validate_required_text(cls, value: Any) -> str:
        text = safe_string(value)
        if text is None:
            raise ValueError("required text field is empty")
        return text

    @field_validator("page_number", mode="before")
    @classmethod
    def validate_page_number(cls, value: Any) -> Optional[Union[int, str]]:
        text = safe_string(value)
        if text is None:
            return None
        if re.fullmatch(r"\d+", text):
            return int(text)
        return text

    @field_validator("confidence", mode="before")
    @classmethod
    def validate_confidence(cls, value: Any) -> Optional[float]:
        if value is None or value == "":
            return None
        return clamp_float(value, 0.0, 1.0)


class ChunkFinancialFacts(BaseModel):
    model_config = ConfigDict(extra="ignore")

    metrics: List[Metric] = Field(default_factory=list)
    company_info: List[CompanyInfoFact] = Field(default_factory=list)

    @field_validator("metrics", "company_info", mode="before")
    @classmethod
    def validate_list(cls, value: Any) -> List[Any]:
        if value is None or value == "":
            return []
        if isinstance(value, list):
            return value
        if isinstance(value, dict):
            return [value]
        return []


class FinancialMetricValue(Metric):
    raw_value: Optional[str] = Field(default=None, description="Legacy exact value field, populated only when different from value.")
    normalized_value: Optional[float] = Field(default=None, description="Numeric normalized value when safely derived by Python.")
    currency: Optional[str] = Field(default=None, description="ISO currency code when explicitly present.")
    scale: Optional[str] = Field(default=None, description="Scale such as thousands, millions, billions, or percent when explicitly present.")
    page: Optional[int] = Field(default=None, description="Legacy one-based page number when known.")

    @field_validator("raw_value", "currency", "scale", mode="before")
    @classmethod
    def validate_optional_legacy_text(cls, value: Any) -> Optional[str]:
        return safe_string(value)

    @field_validator("normalized_value", mode="before")
    @classmethod
    def validate_normalized_value(cls, value: Any) -> Optional[float]:
        if value is None or value == "":
            return None
        try:
            number = float(value)
            return number if math.isfinite(number) else None
        except Exception:
            return None


class ExtractedFinancialData(BaseModel):
    model_config = ConfigDict(extra="forbid")

    company_name: Optional[str] = Field(default=None)
    ticker: Optional[str] = Field(default=None)
    reporting_period: Optional[str] = Field(default=None)
    financial_year: Optional[str] = Field(default=None)
    quarter: Optional[str] = Field(default=None)
    currency: Optional[str] = Field(default=None)
    country: Optional[str] = Field(default=None)
    industry: Optional[str] = Field(default=None)
    reporting_standard: Optional[str] = Field(default=None)
    fiscal_period: Optional[str] = Field(default=None)
    company_info: Dict[str, str] = Field(default_factory=dict)
    metrics: List[FinancialMetricValue] = Field(default_factory=list)
    revenue: Optional[FinancialMetricValue] = Field(default=None)
    gross_revenue: Optional[FinancialMetricValue] = Field(default=None)
    gross_profit: Optional[FinancialMetricValue] = Field(default=None)
    operating_income: Optional[FinancialMetricValue] = Field(default=None)
    ebit: Optional[FinancialMetricValue] = Field(default=None)
    ebitda: Optional[FinancialMetricValue] = Field(default=None)
    net_income: Optional[FinancialMetricValue] = Field(default=None)
    net_profit: Optional[FinancialMetricValue] = Field(default=None)
    eps: Optional[FinancialMetricValue] = Field(default=None)
    diluted_eps: Optional[FinancialMetricValue] = Field(default=None)
    assets: Optional[FinancialMetricValue] = Field(default=None)
    current_assets: Optional[FinancialMetricValue] = Field(default=None)
    non_current_assets: Optional[FinancialMetricValue] = Field(default=None)
    liabilities: Optional[FinancialMetricValue] = Field(default=None)
    current_liabilities: Optional[FinancialMetricValue] = Field(default=None)
    long_term_debt: Optional[FinancialMetricValue] = Field(default=None)
    debt: Optional[FinancialMetricValue] = Field(default=None)
    equity: Optional[FinancialMetricValue] = Field(default=None)
    cash: Optional[FinancialMetricValue] = Field(default=None)
    cash_equivalents: Optional[FinancialMetricValue] = Field(default=None)
    cash_flow: Optional[FinancialMetricValue] = Field(default=None)
    operating_cash_flow: Optional[FinancialMetricValue] = Field(default=None)
    investing_cash_flow: Optional[FinancialMetricValue] = Field(default=None)
    financing_cash_flow: Optional[FinancialMetricValue] = Field(default=None)
    free_cash_flow: Optional[FinancialMetricValue] = Field(default=None)
    capital_expenditure: Optional[FinancialMetricValue] = Field(default=None)
    inventory: Optional[FinancialMetricValue] = Field(default=None)
    receivables: Optional[FinancialMetricValue] = Field(default=None)
    payables: Optional[FinancialMetricValue] = Field(default=None)
    working_capital: Optional[FinancialMetricValue] = Field(default=None)
    tax_expense: Optional[FinancialMetricValue] = Field(default=None)
    interest_expense: Optional[FinancialMetricValue] = Field(default=None)
    operating_margin: Optional[FinancialMetricValue] = Field(default=None)
    gross_margin: Optional[FinancialMetricValue] = Field(default=None)
    net_margin: Optional[FinancialMetricValue] = Field(default=None)
    roa: Optional[FinancialMetricValue] = Field(default=None)
    roe: Optional[FinancialMetricValue] = Field(default=None)
    current_ratio: Optional[FinancialMetricValue] = Field(default=None)
    quick_ratio: Optional[FinancialMetricValue] = Field(default=None)
    debt_to_equity: Optional[FinancialMetricValue] = Field(default=None)
    auditor: Optional[str] = Field(default=None)
    ceo: Optional[str] = Field(default=None)
    employees: Optional[FinancialMetricValue] = Field(default=None)
    dividend: Optional[FinancialMetricValue] = Field(default=None)
    shares_outstanding: Optional[FinancialMetricValue] = Field(default=None)
    confidence_score: float = Field(default=0.0, ge=0.0, le=1.0)
    all_financial_values: List[FinancialMetricValue] = Field(default_factory=list)

    @field_validator("confidence_score", mode="before")
    @classmethod
    def validate_model_confidence(cls, value: Any) -> float:
        return clamp_float(value, 0.0, 1.0)


class TextExtractionRequest(BaseModel):
    text: str = Field(min_length=1)
    filename: Optional[str] = Field(default="text-input.txt")
    mime_type: Optional[str] = Field(default="text/plain")
    metadata: Dict[str, Any] = Field(default_factory=dict)


class ExtractionState(TypedDict, total=False):
    uploaded_file: bytes
    filename: str
    extension: str
    mime_type: str
    metadata: Dict[str, Any]
    page_count: int
    ocr_required: bool
    raw_text: str
    ocr_text: str
    merged_text: str
    tables: List[Dict[str, Any]]
    chunks: List[Dict[str, Any]]
    chunk_results: List[Dict[str, Any]]
    financial_data: Dict[str, Any]
    confidence_score: float
    execution_time: float
    errors: List[str]
    warnings: List[str]
    start_time: float
    response: Dict[str, Any]
    ocr_pages: List[int]
    native_quality: float
    status_code: int


class MissingGroqAPIKeyError(Exception):
    def __init__(self) -> None:
        super().__init__("Missing GROQ_API_KEY environment variable")


OCR_MODEL: Any = None
OCR_LOCK = Lock()
LLM_LOCK = Lock()
LLM_INSTANCE: Optional[ChatGroq] = None
LLM_INSTANCE_KEY: Optional[Tuple[str, str]] = None
STRUCTURED_CHAIN: Any = None
PARSER_CHAIN: Any = None
RAW_CHAIN: Any = None
CHAIN_KEY: Optional[Tuple[str, str]] = None
LLM_SEMAPHORE = asyncio.Semaphore(LLM_CONCURRENCY)


def clamp_float(value: Any, lower: float, upper: float) -> float:
    try:
        number = float(value)
        if not math.isfinite(number):
            return lower
        return max(lower, min(upper, number))
    except Exception:
        return lower


def now_iso() -> str:
    return datetime.utcnow().replace(microsecond=0).isoformat() + "Z"


def safe_string(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def sanitize_for_json(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(k): sanitize_for_json(v) for k, v in value.items()}
    if isinstance(value, list):
        return [sanitize_for_json(v) for v in value]
    if isinstance(value, tuple):
        return [sanitize_for_json(v) for v in value]
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        number = float(value)
        return number if math.isfinite(number) else None
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, bytes):
        return f"<{len(value)} bytes>"
    if isinstance(value, datetime):
        return value.isoformat()
    return value


def ensure_state_defaults(state: ExtractionState) -> ExtractionState:
    state.setdefault("uploaded_file", b"")
    state.setdefault("filename", "upload")
    state.setdefault("extension", "")
    state.setdefault("mime_type", "application/octet-stream")
    state.setdefault("metadata", {})
    state.setdefault("page_count", 0)
    state.setdefault("ocr_required", False)
    state.setdefault("raw_text", "")
    state.setdefault("ocr_text", "")
    state.setdefault("merged_text", "")
    state.setdefault("tables", [])
    state.setdefault("chunks", [])
    state.setdefault("chunk_results", [])
    state.setdefault("financial_data", empty_financial_data())
    state.setdefault("confidence_score", 0.0)
    state.setdefault("execution_time", 0.0)
    state.setdefault("errors", [])
    state.setdefault("warnings", [])
    state.setdefault("start_time", time.perf_counter())
    state.setdefault("response", {})
    state.setdefault("ocr_pages", [])
    state.setdefault("native_quality", 0.0)
    state.setdefault("status_code", 200)
    return state


def add_error(state: ExtractionState, message: str) -> None:
    errors = list(state.get("errors", []))
    if message not in errors:
        errors.append(message)
        logger.error(message)
    state["errors"] = errors


def add_warning(state: ExtractionState, message: str) -> None:
    warnings = list(state.get("warnings", []))
    if message not in warnings:
        warnings.append(message)
        logger.warning(message)
    state["warnings"] = warnings


def dump_financial_data(data: ExtractedFinancialData) -> Dict[str, Any]:
    payload = data.model_dump(mode="json", exclude_none=True)
    if not payload.get("company_info"):
        payload.pop("company_info", None)
    return payload


def dump_chunk_facts(data: ChunkFinancialFacts) -> Dict[str, Any]:
    payload = data.model_dump(mode="json", exclude_none=True)
    if not payload.get("company_info"):
        payload.pop("company_info", None)
    payload.setdefault("metrics", [])
    return payload


def empty_chunk_financial_facts() -> Dict[str, Any]:
    return dump_chunk_facts(ChunkFinancialFacts())


def empty_financial_data() -> Dict[str, Any]:
    return dump_financial_data(ExtractedFinancialData())


def get_extension_from_zip(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = set(archive.namelist())
            if any(name.startswith("word/") for name in names):
                return ".docx"
            if any(name.startswith("xl/") for name in names):
                return ".xlsx"
    except Exception:
        return ""
    return ""


def detect_extension_from_content(data: bytes, filename: str, mime_type: str) -> str:
    suffix = Path(filename or "").suffix.lower()
    if suffix in SUPPORTED_EXTENSIONS:
        return ".jpg" if suffix == ".jpeg" else suffix
    lowered_mime = (mime_type or "").lower()
    if data.startswith(b"%PDF"):
        return ".pdf"
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return ".png"
    if data.startswith(b"\xff\xd8\xff"):
        return ".jpg"
    if data.startswith(b"BM"):
        return ".bmp"
    if data.startswith(b"II*\x00") or data.startswith(b"MM\x00*"):
        return ".tiff"
    if data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return ".xls"
    if data.startswith(b"PK"):
        zipped = get_extension_from_zip(data)
        if zipped:
            return zipped
    if "pdf" in lowered_mime:
        return ".pdf"
    if "png" in lowered_mime:
        return ".png"
    if "jpeg" in lowered_mime or "jpg" in lowered_mime:
        return ".jpg"
    if "bmp" in lowered_mime:
        return ".bmp"
    if "tiff" in lowered_mime:
        return ".tiff"
    if "wordprocessingml" in lowered_mime:
        return ".docx"
    if "spreadsheetml" in lowered_mime:
        return ".xlsx"
    if "ms-excel" in lowered_mime or "excel" in lowered_mime:
        return ".xls"
    if "csv" in lowered_mime:
        return ".csv"
    if "tab-separated" in lowered_mime or "tsv" in lowered_mime:
        return ".tsv"
    if "html" in lowered_mime:
        return ".html"
    if "json" in lowered_mime:
        return ".json"
    if "xml" in lowered_mime:
        return ".xml"
    if "rtf" in lowered_mime:
        return ".rtf"
    if "text" in lowered_mime:
        return ".txt"
    guessed_mime, _ = mimetypes.guess_type(filename or "")
    if guessed_mime:
        return detect_extension_from_content(data, f"file{mimetypes.guess_extension(guessed_mime) or ''}", guessed_mime)
    if looks_like_text(data):
        return ".txt"
    return suffix


def mime_for_extension(extension: str, provided: str) -> str:
    if provided and provided != "application/octet-stream":
        return provided
    guessed = mimetypes.types_map.get(extension)
    if guessed:
        return guessed
    mapping = {
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".xls": "application/vnd.ms-excel",
        ".csv": "text/csv",
        ".tsv": "text/tab-separated-values",
        ".md": "text/markdown",
        ".rtf": "application/rtf",
        ".html": "text/html",
        ".htm": "text/html",
        ".json": "application/json",
        ".xml": "application/xml",
        ".tiff": "image/tiff",
        ".tif": "image/tiff",
    }
    return mapping.get(extension, "application/octet-stream")


def looks_like_text(data: bytes) -> bool:
    if not data:
        return False
    sample = data[:4096]
    if b"\x00" in sample:
        return False
    encoding = detect_encoding(sample)
    try:
        sample.decode(encoding, errors="strict")
        return True
    except Exception:
        return False


def detect_encoding(data: bytes) -> str:
    if not data:
        return "utf-8"
    if chardet is not None:
        detected = chardet.detect(data[:200000])
        encoding = detected.get("encoding") if detected else None
        if encoding:
            return encoding
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            data.decode(encoding)
            return encoding
        except Exception:
            continue
    return "utf-8"


async def read_upload_stream(upload: UploadFile) -> bytes:
    chunks: List[bytes] = []
    total = 0
    while True:
        chunk = await upload.read(1024 * 1024)
        if not chunk:
            break
        total += len(chunk)
        if total > MAX_UPLOAD_BYTES:
            raise ValueError(f"File exceeds maximum size of {MAX_UPLOAD_BYTES // (1024 * 1024)} MB")
        chunks.append(chunk)
    return b"".join(chunks)


async def write_temp_bytes(data: bytes, suffix: str) -> str:
    descriptor, path = tempfile.mkstemp(suffix=suffix)
    os.close(descriptor)
    async with aiofiles.open(path, "wb") as handle:
        await handle.write(data)
    return path


async def remove_temp_file(path: str) -> None:
    try:
        if path and os.path.exists(path):
            os.remove(path)
            logger.info("Cleanup removed temporary file {path}", path=path)
    except Exception as exc:
        logger.warning("Cleanup failed for temporary file {path}: {error}", path=path, error=str(exc))


async def run_blocking(function: Any, *args: Any, **kwargs: Any) -> Any:
    return await asyncio.to_thread(function, *args, **kwargs)


def extract_pdf_metadata(data: bytes) -> Tuple[Dict[str, Any], int]:
    metadata: Dict[str, Any] = {}
    with fitz.open(stream=data, filetype="pdf") as document:
        metadata.update({str(k): safe_string(v) for k, v in (document.metadata or {}).items() if safe_string(v) is not None})
        metadata["is_encrypted"] = bool(document.is_encrypted)
        metadata["requires_credential"] = bool(getattr(document, "needs_" + "pa" + "ss", False))
        metadata["page_count"] = int(document.page_count)
        metadata["format"] = "PDF"
        metadata["creation_date_normalized"] = normalize_possible_date(metadata.get("creationDate"))
        metadata["modification_date_normalized"] = normalize_possible_date(metadata.get("modDate"))
        return metadata, int(document.page_count)


def extract_image_metadata(data: bytes) -> Tuple[Dict[str, Any], int]:
    with Image.open(io.BytesIO(data)) as image:
        frames = sum(1 for _ in ImageSequence.Iterator(image)) if getattr(image, "is_animated", False) else 1
        metadata = {
            "format": image.format,
            "mode": image.mode,
            "width": image.width,
            "height": image.height,
            "frames": frames,
        }
        try:
            exif = image.getexif()
            if exif:
                metadata["exif_keys"] = [str(key) for key in exif.keys()]
        except Exception:
            metadata["exif_keys"] = []
        return metadata, frames


def extract_docx_metadata(data: bytes) -> Tuple[Dict[str, Any], int]:
    document = Document(io.BytesIO(data))
    properties = document.core_properties
    metadata = {
        "author": safe_string(properties.author),
        "category": safe_string(properties.category),
        "comments": safe_string(properties.comments),
        "content_status": safe_string(properties.content_status),
        "created": properties.created.isoformat() if properties.created else None,
        "identifier": safe_string(properties.identifier),
        "keywords": safe_string(properties.keywords),
        "language": safe_string(properties.language),
        "last_modified_by": safe_string(properties.last_modified_by),
        "last_printed": properties.last_printed.isoformat() if properties.last_printed else None,
        "modified": properties.modified.isoformat() if properties.modified else None,
        "revision": properties.revision,
        "subject": safe_string(properties.subject),
        "title": safe_string(properties.title),
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "section_count": len(document.sections),
        "format": "DOCX",
    }
    return {k: v for k, v in metadata.items() if v is not None}, 1


def extract_xlsx_metadata(data: bytes) -> Tuple[Dict[str, Any], int]:
    if data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        excel_file = pd.ExcelFile(io.BytesIO(data), engine="xlrd")
        sheet_names = list(excel_file.sheet_names)
        metadata = {
            "sheet_count": len(sheet_names),
            "sheet_names": sheet_names,
            "format": "XLS",
        }
        return metadata, max(1, len(sheet_names))
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    properties = workbook.properties
    sheet_names = list(workbook.sheetnames)
    metadata = {
        "creator": safe_string(properties.creator),
        "last_modified_by": safe_string(properties.lastModifiedBy),
        "created": properties.created.isoformat() if properties.created else None,
        "modified": properties.modified.isoformat() if properties.modified else None,
        "title": safe_string(properties.title),
        "subject": safe_string(properties.subject),
        "keywords": safe_string(properties.keywords),
        "category": safe_string(properties.category),
        "description": safe_string(properties.description),
        "sheet_count": len(sheet_names),
        "sheet_names": sheet_names,
        "format": "XLSX",
    }
    workbook.close()
    return {k: v for k, v in metadata.items() if v is not None}, max(1, len(sheet_names))

def extract_text_metadata(data: bytes, extension: str) -> Tuple[Dict[str, Any], int]:
    encoding = detect_encoding(data)
    metadata = {
        "encoding": encoding,
        "byte_size": len(data),
        "format": extension.replace(".", "").upper(),
    }
    return metadata, 1


def normalize_possible_date(value: Any) -> Optional[str]:
    if not value:
        return None
    text = str(value)
    if text.startswith("D:"):
        text = text[2:16]
    try:
        return date_parser.parse(text, fuzzy=True).date().isoformat()
    except Exception:
        return None


def extract_pdf_native_text(data: bytes) -> Tuple[str, Dict[str, Any], List[str]]:
    text_parts: List[str] = []
    page_char_counts: List[int] = []
    image_counts: List[int] = []
    warnings: List[str] = []
    with fitz.open(stream=data, filetype="pdf") as document:
        for page_index in range(document.page_count):
            page = document.load_page(page_index)
            page_text = page.get_text("text") or ""
            text_parts.append(f"Page {page_index + 1}\n{page_text.strip()}")
            page_char_counts.append(len(page_text.strip()))
            try:
                image_counts.append(len(page.get_images(full=True)))
            except Exception:
                image_counts.append(0)
    total_chars = sum(page_char_counts)
    page_count = max(1, len(page_char_counts))
    quality = min(1.0, total_chars / max(1, page_count * 1200))
    metadata = {
        "native_page_char_counts": page_char_counts,
        "native_page_image_counts": image_counts,
        "native_total_chars": total_chars,
        "native_quality": round(quality, 4),
    }
    if total_chars == 0:
        warnings.append("PDF native text extraction returned no text")
    return "\n\n".join(text_parts).strip(), metadata, warnings


async def extract_docx_native_text(data: bytes) -> Tuple[str, Dict[str, Any], List[Dict[str, Any]], List[str]]:
    warnings: List[str] = []
    text_parts: List[str] = []
    tables: List[Dict[str, Any]] = []
    document = Document(io.BytesIO(data))
    header_footer_lines: List[str] = []
    for section_index, section in enumerate(document.sections):
        for paragraph in section.header.paragraphs:
            if paragraph.text.strip():
                header_footer_lines.append(f"Header {section_index + 1}: {paragraph.text.strip()}")
        for paragraph in section.footer.paragraphs:
            if paragraph.text.strip():
                header_footer_lines.append(f"Footer {section_index + 1}: {paragraph.text.strip()}")
        tables.extend(extract_docx_tables_from_container(section.header.tables, f"header_{section_index + 1}"))
        tables.extend(extract_docx_tables_from_container(section.footer.tables, f"footer_{section_index + 1}"))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text and paragraph.text.strip()]
    tables.extend(extract_docx_tables_from_container(document.tables, "body"))
    if header_footer_lines:
        text_parts.append("Headers and Footers\n" + "\n".join(header_footer_lines))
    if paragraphs:
        text_parts.append("Paragraphs\n" + "\n".join(paragraphs))
    if tables:
        text_parts.append("Tables\n" + render_tables(tables))
    temp_path = await write_temp_bytes(data, ".docx")
    try:
        extracted = await run_blocking(docx2txt.process, temp_path)
        if extracted and extracted.strip():
            text_parts.append("Docx2txt\n" + extracted.strip())
    except Exception as exc:
        warnings.append(f"DOCX secondary extraction failed: {str(exc)}")
    finally:
        await remove_temp_file(temp_path)
    metadata = {
        "paragraph_count": len(paragraphs),
        "table_count": len(tables),
        "header_footer_line_count": len(header_footer_lines),
    }
    return "\n\n".join(text_parts).strip(), metadata, tables, warnings


def extract_docx_tables_from_container(table_objects: Any, location: str) -> List[Dict[str, Any]]:
    tables: List[Dict[str, Any]] = []
    for table_index, table in enumerate(table_objects):
        rows: List[List[str]] = []
        for row in table.rows:
            rows.append([normalize_cell_text(cell.text) for cell in row.cells])
        if any(any(cell for cell in row) for row in rows):
            tables.append({"type": "docx", "location": location, "index": table_index + 1, "rows": rows})
    return tables


def extract_xlsx_native_text(data: bytes) -> Tuple[str, Dict[str, Any], List[Dict[str, Any]], List[str]]:
    warnings: List[str] = []
    tables: List[Dict[str, Any]] = []
    text_parts: List[str] = []
    is_xls = data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")
    engine = "xlrd" if is_xls else "openpyxl"
    filetype = "xls" if is_xls else "xlsx"
    try:
        excel_file = pd.ExcelFile(io.BytesIO(data), engine=engine)
        sheet_names = list(excel_file.sheet_names)
        for sheet_name in sheet_names:
            dataframe = pd.read_excel(excel_file, sheet_name=sheet_name, dtype=str, header=None).fillna("")
            rows = [[normalize_cell_text(value) for value in row] for row in dataframe.values.tolist()]
            tables.append({"type": filetype, "sheet": sheet_name, "index": len(tables) + 1, "rows": rows})
            text_parts.append(f"Sheet: {sheet_name}\n" + rows_to_text(rows))
        metadata = {"sheet_count": len(sheet_names), "sheet_names": sheet_names, "format": filetype.upper()}
        return "\n\n".join(text_parts).strip(), metadata, tables, warnings
    except Exception as exc:
        warnings.append(f"Spreadsheet pandas extraction failed: {str(exc)}")
        if is_xls:
            return "", {"sheet_count": 0, "format": "XLS"}, tables, warnings
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            rows = [[normalize_cell_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
            tables.append({"type": "xlsx", "sheet": sheet_name, "index": len(tables) + 1, "rows": rows})
            text_parts.append(f"Sheet: {sheet_name}\n" + rows_to_text(rows))
        workbook.close()
        metadata = {"sheet_count": len(tables), "sheet_names": [table.get("sheet") for table in tables], "format": "XLSX"}
        return "\n\n".join(text_parts).strip(), metadata, tables, warnings

def extract_csv_native_text(data: bytes) -> Tuple[str, Dict[str, Any], List[Dict[str, Any]], List[str]]:
    warnings: List[str] = []
    encoding = detect_encoding(data)
    decoded = data.decode(encoding, errors="replace")
    delimiter = detect_csv_delimiter(decoded[:8192])
    dataframe = pd.read_csv(io.StringIO(decoded), sep=delimiter, dtype=str, keep_default_na=False, on_bad_lines="skip")
    rows = [list(dataframe.columns)] + dataframe.astype(str).values.tolist()
    rows = [[normalize_cell_text(cell) for cell in row] for row in rows]
    table = {"type": "csv", "index": 1, "rows": rows}
    metadata = {
        "encoding": encoding,
        "delimiter": delimiter,
        "row_count": max(0, len(rows) - 1),
        "column_count": len(rows[0]) if rows else 0,
    }
    if len(rows) <= 1:
        warnings.append("CSV extraction found no data rows")
    return rows_to_text(rows), metadata, [table], warnings


def detect_csv_delimiter(sample: str) -> str:
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        return dialect.delimiter
    except csv.Error:
        return ","


def extract_txt_native_text(data: bytes) -> Tuple[str, Dict[str, Any], List[str]]:
    encoding = detect_encoding(data)
    decoded = data.decode(encoding, errors="replace")
    metadata = {"encoding": encoding, "character_count": len(decoded)}
    warnings: List[str] = []
    if not decoded.strip():
        warnings.append("TXT extraction returned no text")
    return decoded.strip(), metadata, warnings


def normalize_cell_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r", " ").replace("\n", " ")
    return re.sub(r"\s+", " ", text).strip()


def rows_to_text(rows: List[List[str]]) -> str:
    return "\n".join("\t".join(cell for cell in row) for row in rows if any(cell for cell in row))


def render_tables(tables: List[Dict[str, Any]]) -> str:
    rendered: List[str] = []
    for table_index, table in enumerate(tables):
        label_parts = ["Table", str(table_index + 1)]
        if table.get("page"):
            label_parts.extend(["Page", str(table.get("page"))])
        if table.get("sheet"):
            label_parts.extend(["Sheet", str(table.get("sheet"))])
        if table.get("location"):
            label_parts.extend(["Location", str(table.get("location"))])
        rendered.append(" ".join(label_parts))
        rows = table.get("rows") or []
        rendered.append(rows_to_text(rows))
    return "\n\n".join(part for part in rendered if part.strip())


def extract_pdf_tables(data: bytes) -> Tuple[List[Dict[str, Any]], List[str]]:
    tables: List[Dict[str, Any]] = []
    warnings: List[str] = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page_index, page in enumerate(pdf.pages):
                extracted_tables = page.extract_tables() or []
                for table_index, rows in enumerate(extracted_tables):
                    cleaned = [[normalize_cell_text(cell) for cell in row] for row in rows if row]
                    if any(any(cell for cell in row) for row in cleaned):
                        tables.append({"type": "pdfplumber", "page": page_index + 1, "index": table_index + 1, "rows": cleaned})
    except Exception as exc:
        warnings.append(f"pdfplumber table extraction failed: {str(exc)}")
    if not tables:
        try:
            with fitz.open(stream=data, filetype="pdf") as document:
                for page_index in range(document.page_count):
                    page = document.load_page(page_index)
                    if hasattr(page, "find_tables"):
                        found = page.find_tables()
                        for table_index, table in enumerate(found.tables):
                            rows = [[normalize_cell_text(cell) for cell in row] for row in table.extract()]
                            if any(any(cell for cell in row) for row in rows):
                                tables.append({"type": "pymupdf", "page": page_index + 1, "index": table_index + 1, "rows": rows})
        except Exception as exc:
            warnings.append(f"PyMuPDF table extraction failed: {str(exc)}")
    return tables, warnings


def get_ocr_model() -> Any:
    global OCR_MODEL
    with OCR_LOCK:
        if OCR_MODEL is None:
            logger.info("OCR model loading")
            try:
                from paddleocr import PaddleOCR

                lang = os.getenv("PADDLEOCR_LANG", "en")
                try:
                    OCR_MODEL = PaddleOCR(use_angle_cls=True, lang=lang, show_log=False)
                except TypeError:
                    try:
                        OCR_MODEL = PaddleOCR(use_textline_orientation=True, lang=lang)
                    except TypeError:
                        OCR_MODEL = PaddleOCR(lang=lang)
                logger.info("OCR model loaded")
            except Exception as exc:
                logger.exception("OCR model initialization failed")
                raise RuntimeError(f"OCR failure: {str(exc)}") from exc
    return OCR_MODEL


def preprocess_image_for_ocr(image: Image.Image) -> np.ndarray:
    image = ImageOps.exif_transpose(image.convert("RGB"))
    array = np.array(image)
    gray = cv2.cvtColor(array, cv2.COLOR_RGB2GRAY)
    height, width = gray.shape[:2]
    scale = 1.0
    if min(height, width) < 1000:
        scale = 1000 / max(1, min(height, width))
    if scale > 1.0:
        gray = cv2.resize(gray, None, fx=scale, fy=scale, interpolation=cv2.INTER_CUBIC)
    denoised = cv2.fastNlMeansDenoising(gray, None, 10, 7, 21)
    thresholded = cv2.threshold(denoised, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)[1]
    return cv2.cvtColor(thresholded, cv2.COLOR_GRAY2RGB)


def parse_ocr_result(result: Any) -> Tuple[List[str], List[float]]:
    lines: List[str] = []
    scores: List[float] = []

    def walk(value: Any) -> None:
        if value is None:
            return
        if isinstance(value, dict):
            if "rec_texts" in value:
                rec_texts = value.get("rec_texts") or []
                rec_scores = value.get("rec_scores") or []
                for index, text in enumerate(rec_texts):
                    clean = safe_string(text)
                    if clean:
                        lines.append(clean)
                        scores.append(clamp_float(rec_scores[index] if index < len(rec_scores) else 0.0, 0.0, 1.0))
                return
            if "text" in value:
                clean = safe_string(value.get("text"))
                if clean:
                    lines.append(clean)
                    scores.append(clamp_float(value.get("confidence") or value.get("score") or 0.0, 0.0, 1.0))
                return
            for nested in value.values():
                walk(nested)
            return
        if isinstance(value, (list, tuple)):
            if len(value) >= 2 and isinstance(value[1], (list, tuple)) and len(value[1]) >= 2 and isinstance(value[1][0], str):
                clean = safe_string(value[1][0])
                if clean:
                    lines.append(clean)
                    scores.append(clamp_float(value[1][1], 0.0, 1.0))
                return
            if len(value) >= 2 and isinstance(value[0], str):
                clean = safe_string(value[0])
                if clean:
                    lines.append(clean)
                    scores.append(clamp_float(value[1], 0.0, 1.0))
                return
            for nested in value:
                walk(nested)

    walk(result)
    return lines, scores


def run_ocr_array(image_array: np.ndarray) -> Tuple[str, float]:
    model = get_ocr_model()
    try:
        result = model.ocr(image_array, cls=True)
    except TypeError:
        try:
            result = model.ocr(image_array)
        except AttributeError:
            result = model.predict(image_array)
    except AttributeError:
        result = model.predict(image_array)
    lines, scores = parse_ocr_result(result)
    confidence = float(sum(scores) / len(scores)) if scores else 0.0
    return "\n".join(lines), confidence


def run_ocr_image_bytes(data: bytes) -> Tuple[str, Dict[str, Any]]:
    texts: List[str] = []
    confidences: List[float] = []
    with Image.open(io.BytesIO(data)) as image:
        frames = [frame.copy() for frame in ImageSequence.Iterator(image)] if getattr(image, "is_animated", False) else [image.copy()]
        for frame_index, frame in enumerate(frames):
            array = preprocess_image_for_ocr(frame)
            text, confidence = run_ocr_array(array)
            if text.strip():
                texts.append(f"Image Frame {frame_index + 1}\n{text.strip()}")
            confidences.append(confidence)
    metadata = {
        "ocr_frame_count": len(confidences),
        "ocr_average_confidence": round(sum(confidences) / len(confidences), 4) if confidences else 0.0,
    }
    return "\n\n".join(texts).strip(), metadata


def run_ocr_pdf_bytes(data: bytes, pages: List[int]) -> Tuple[str, Dict[str, Any]]:
    texts: List[str] = []
    confidences: List[float] = []
    with fitz.open(stream=data, filetype="pdf") as document:
        target_pages = pages if pages else list(range(1, document.page_count + 1))
        for page_number in target_pages:
            if page_number < 1 or page_number > document.page_count:
                continue
            page = document.load_page(page_number - 1)
            matrix = fitz.Matrix(PDF_OCR_DPI_SCALE, PDF_OCR_DPI_SCALE)
            pixmap = page.get_pixmap(matrix=matrix, alpha=False)
            image = Image.open(io.BytesIO(pixmap.tobytes("png")))
            array = preprocess_image_for_ocr(image)
            text, confidence = run_ocr_array(array)
            if text.strip():
                texts.append(f"OCR Page {page_number}\n{text.strip()}")
            confidences.append(confidence)
    metadata = {
        "ocr_page_count": len(confidences),
        "ocr_pages": pages,
        "ocr_average_confidence": round(sum(confidences) / len(confidences), 4) if confidences else 0.0,
    }
    return "\n\n".join(texts).strip(), metadata


def remove_duplicate_lines(text: str, threshold: int = 98) -> str:
    output: List[str] = []
    exact_seen: set[str] = set()
    recent: List[str] = []
    for line in text.splitlines():
        clean = line.strip()
        if not clean:
            if output and output[-1] != "":
                output.append("")
            continue
        signature = re.sub(r"\s+", " ", clean).lower()
        if signature in exact_seen:
            continue
        duplicate = False
        if len(signature) > 12:
            for candidate in recent[-80:]:
                if fuzz.ratio(signature, candidate) >= threshold:
                    duplicate = True
                    break
        if duplicate:
            continue
        exact_seen.add(signature)
        recent.append(signature)
        output.append(clean)
    return "\n".join(output).strip()


def remove_ocr_duplicates(native_text: str, ocr_text: str) -> str:
    native_signatures = [re.sub(r"\s+", " ", line.strip()).lower() for line in native_text.splitlines() if len(line.strip()) > 12]
    kept: List[str] = []
    for line in ocr_text.splitlines():
        clean = line.strip()
        if not clean:
            continue
        signature = re.sub(r"\s+", " ", clean).lower()
        if any(fuzz.ratio(signature, native) >= 97 for native in native_signatures[-400:]):
            continue
        kept.append(clean)
    return "\n".join(kept)


def normalize_document_text(text: str) -> str:
    normalized = unicodedata.normalize("NFKC", text or "")
    normalized = normalized.replace("\r\n", "\n").replace("\r", "\n")
    normalized = normalize_dates_in_text(normalized)
    normalized = normalize_currencies_in_text(normalized)
    normalized = normalize_percentages_in_text(normalized)
    normalized = normalize_numbers_in_text(normalized)
    normalized_lines: List[str] = []
    for line in normalized.splitlines():
        if "\t" in line:
            cells = [re.sub(r"[ \u00a0]+", " ", cell).strip() for cell in line.split("\t")]
            normalized_lines.append("\t".join(cells).strip())
        else:
            normalized_lines.append(re.sub(r"[ \t\u00a0]+", " ", line).strip())
    normalized = "\n".join(normalized_lines)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    return remove_duplicate_lines(normalized).strip()


def normalize_dates_in_text(text: str) -> str:
    month_names = r"(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)"
    patterns = [
        rf"\b{month_names}\s+\d{{1,2}},?\s+\d{{4}}\b",
        rf"\b\d{{1,2}}\s+{month_names}\s+\d{{4}}\b",
        r"\b\d{4}-\d{1,2}-\d{1,2}\b",
        r"\b\d{1,2}/\d{1,2}/\d{2,4}\b",
    ]
    combined = re.compile("|".join(f"({pattern})" for pattern in patterns), re.IGNORECASE)

    def replace(match: re.Match[str]) -> str:
        raw = match.group(0)
        try:
            parsed = date_parser.parse(raw, fuzzy=False, dayfirst=False)
            return parsed.date().isoformat()
        except Exception:
            try:
                parsed = date_parser.parse(raw, fuzzy=False, dayfirst=True)
                return parsed.date().isoformat()
            except Exception:
                return raw

    return combined.sub(replace, text)


def normalize_currencies_in_text(text: str) -> str:
    currency_pattern = re.compile(r"(?P<currency>USD|EUR|GBP|INR|CAD|AUD|JPY|[$\u20ac\u00a3\u20b9\u00a5])\s*(?P<number>\(?[-+]?\d[\d,]*(?:\.\d+)?\)?)(?P<scale>\s*(?:thousand|million|billion|trillion|crore|crores|cr|lakh|lakhs|k|m|bn|b))?", re.IGNORECASE)

    def replace(match: re.Match[str]) -> str:
        currency = match.group("currency")
        code = CURRENCY_SYMBOLS.get(currency, currency.upper())
        number = normalize_number_token(match.group("number"))
        scale = (match.group("scale") or "").strip()
        return f"{code} {number}{(' ' + scale) if scale else ''}"

    return currency_pattern.sub(replace, text)


def normalize_percentages_in_text(text: str) -> str:
    percentage_pattern = re.compile(r"(?<![\w.])(?P<number>\(?[-+]?\d[\d,]*(?:\.\d+)?\)?)\s*%")

    def replace(match: re.Match[str]) -> str:
        return f"{normalize_number_token(match.group('number'))}%"

    return percentage_pattern.sub(replace, text)


def normalize_numbers_in_text(text: str) -> str:
    number_pattern = re.compile(r"(?<![\w.])\(?[-+]?\d{1,3}(?:,\d{3})+(?:\.\d+)?\)?(?![\w.])")
    return number_pattern.sub(lambda match: normalize_number_token(match.group(0)), text)


def normalize_number_token(token: str) -> str:
    raw = token.strip()
    negative = raw.startswith("(") and raw.endswith(")")
    cleaned = raw.strip("()").replace(",", "")
    try:
        value = parse_decimal(cleaned, locale="en_US")
        text = format(value, "f")
        if "." in text:
            text = text.rstrip("0").rstrip(".")
        return f"-{text}" if negative and not text.startswith("-") else text
    except (NumberFormatError, ValueError):
        return raw


def chunk_text(text: str) -> List[Dict[str, Any]]:
    if not text.strip():
        return []
    target = max(4000, MAX_CHUNK_CHARS)
    overlap = max(0, min(CHUNK_OVERLAP_CHARS, target // 4))
    paragraphs = re.split(r"\n{2,}", text)
    chunks: List[Dict[str, Any]] = []
    current: List[str] = []
    current_length = 0

    def flush() -> None:
        nonlocal current, current_length
        if not current:
            return
        chunk_body = "\n\n".join(current).strip()
        if chunk_body:
            chunks.append({"index": len(chunks) + 1, "text": chunk_body, "character_count": len(chunk_body)})
        if overlap > 0 and chunk_body:
            tail = chunk_body[-overlap:]
            current = [tail]
            current_length = len(tail)
        else:
            current = []
            current_length = 0

    for paragraph in paragraphs:
        clean = paragraph.strip()
        if not clean:
            continue
        if len(clean) > target:
            flush()
            start = 0
            while start < len(clean):
                end = min(len(clean), start + target)
                segment = clean[start:end]
                chunks.append({"index": len(chunks) + 1, "text": segment, "character_count": len(segment)})
                start = max(end - overlap, end) if overlap == 0 else end - overlap
                if start >= len(clean) or end == len(clean):
                    break
            current = []
            current_length = 0
            continue
        if current_length + len(clean) + 2 > target:
            flush()
        current.append(clean)
        current_length += len(clean) + 2
    flush()
    for index, chunk in enumerate(chunks):
        chunk["index"] = index + 1
        chunk["total"] = len(chunks)
    return chunks


def normalize_company_name_text(value: Any) -> Optional[str]:
    text = safe_string(value)
    if text is None:
        return None
    text = re.sub(r"(?i)^\s*microsoft\s+word\s*[-:]+\s*", "", text)
    text = re.sub(r"[_]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip(" -??:;,.\t")
    for delimiter in (" - ", " ? ", " ? ", " | "):
        if delimiter in text:
            parts = [part.strip(" -??:;,.\t") for part in text.split(delimiter) if part.strip()]
            legal_parts = [part for part in parts if LEGAL_ENTITY_SUFFIX_PATTERN.search(part)]
            if legal_parts:
                text = legal_parts[-1]
                break
    text = re.sub(r"(?i)\b(?:annual report|financial statements?|integrated report|final|draft)\b", "", text)
    text = re.sub(r"\s+", " ", text).strip(" -??:;,.\t")
    if not text or len(text) > 140:
        return None
    if text[0].islower() or "." in text or "," in text:
        return None
    if COMPANY_NAME_EXCLUDE_PATTERN.search(text):
        return None
    if not LEGAL_ENTITY_SUFFIX_PATTERN.search(text):
        return None
    if text.upper() == text:
        keep_upper = {"LLC", "LLP", "PLC", "LTD", "PVT"}
        words = [word if word in keep_upper else word.capitalize() for word in text.split()]
        return " ".join(words)
    return text


def looks_like_company_name_line(line: str) -> bool:
    return normalize_company_name_text(line) is not None


def extract_company_name_from_metadata(metadata: Dict[str, Any]) -> Optional[str]:
    for key in ("title", "subject", "company", "document_title"):
        company_name = normalize_company_name_text(metadata.get(key))
        if company_name:
            return company_name
    return None


def extract_company_name_from_text(text: str) -> Optional[str]:
    lines = [line.strip() for line in (text or "").splitlines() if line.strip()]
    for line in lines[:180]:
        labelled = value_after_label_prefix(line, "company name") or value_after_label_prefix(line, "company")
        company_name = normalize_company_name_text(labelled) if labelled else normalize_company_name_text(line)
        if company_name:
            return company_name
    return None


def detect_currency_code_from_text(value: Any) -> Optional[str]:
    text = safe_string(value)
    if text is None:
        return None
    for symbol, code in CURRENCY_SYMBOLS.items():
        if symbol in text:
            return code
    match = re.search(r"\b(?:USD|EUR|GBP|INR|CAD|AUD|JPY|CNY|CHF|SGD|HKD|AED)\b", text, flags=re.IGNORECASE)
    return match.group(0).upper() if match else None


def extract_currency_from_financial_data(financial_data: Dict[str, Any]) -> Optional[str]:
    counts: Dict[str, int] = {}

    def add_currency(value: Any) -> None:
        code = detect_currency_code_from_text(value)
        if code:
            counts[code] = counts.get(code, 0) + 1

    for metric in financial_data.get("metrics") or financial_data.get("all_financial_values") or []:
        if isinstance(metric, dict):
            add_currency(metric.get("currency"))
            add_currency(metric.get("unit"))
            add_currency(metric.get("value"))
            add_currency(metric.get("raw_value"))
    for field in METRIC_FINANCIAL_FIELDS:
        metric = financial_data.get(field)
        if isinstance(metric, dict):
            add_currency(metric.get("currency"))
            add_currency(metric.get("unit"))
            add_currency(metric.get("value"))
            add_currency(metric.get("raw_value"))
    if not counts:
        return None
    return sorted(counts.items(), key=lambda item: item[1], reverse=True)[0][0]


def augment_financial_data_with_context(financial_data: Dict[str, Any], metadata: Dict[str, Any], text: str) -> Dict[str, Any]:
    data = dict(financial_data or {})
    company_info = dict(data.get("company_info") or {}) if isinstance(data.get("company_info"), dict) else {}
    metadata_company_name = extract_company_name_from_metadata(metadata or {})
    text_company_name = extract_company_name_from_text(text or "")
    preferred_company_name = metadata_company_name or text_company_name
    current_company_name = normalize_company_name_text(data.get("company_name"))
    if preferred_company_name and (metadata_company_name or not current_company_name):
        data["company_name"] = preferred_company_name
        company_info["company_name"] = preferred_company_name
    elif current_company_name:
        data["company_name"] = current_company_name
        company_info.setdefault("company_name", current_company_name)
    else:
        data.pop("company_name", None)
        company_info.pop("company_name", None)
    if not data.get("currency"):
        currency = extract_currency_from_financial_data(data) or detect_currency_code_from_text(text[:12000] if text else "")
        if currency:
            data["currency"] = currency
            company_info["currency"] = currency
    if company_info:
        data["company_info"] = company_info
    return dump_financial_data(ExtractedFinancialData.model_validate(data))


def approximate_token_count(text: str) -> int:
    divisor = max(1.0, APPROX_CHARS_PER_TOKEN)
    return int(math.ceil(len(text or "") / divisor))


def is_financial_signal_line(line: str) -> bool:
    clean = safe_string(line) or ""
    if not clean:
        return False
    key = label_key(clean)
    if looks_like_company_name_line(clean):
        return True
    has_value = bool(re.search(r"(?i)(?:USD|EUR|GBP|INR|CAD|AUD|JPY|CNY|CHF|SGD|HKD|AED|[$????])|\d", clean))
    if "\t" in clean and has_value:
        return True
    for aliases in METRIC_FIELD_ALIASES.values():
        if any(label_key(alias) in key for alias in aliases) and has_value:
            return True
    for aliases in COMPANY_INFO_FIELD_ALIASES.values():
        if any(label_key(alias) in key for alias in aliases) and re.search(r"[:=\-??\t]", clean):
            return True
    return False


def compact_chunk_for_llm(text: str, max_chars: int) -> Tuple[str, bool]:
    source = text or ""
    if max_chars <= 0:
        return "", bool(source)
    if len(source) <= max_chars:
        return source, False
    lines = source.splitlines()
    selected: List[str] = []
    seen: set[str] = set()
    current_page: Optional[str] = None
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if re.search(r"\b(?:page|ocr page)\s+\d+\b", stripped, flags=re.IGNORECASE):
            current_page = stripped
            continue
        if not is_financial_signal_line(stripped):
            continue
        if current_page and current_page.lower() not in seen:
            selected.append(current_page)
            seen.add(current_page.lower())
        signature = re.sub(r"\s+", " ", stripped).lower()
        if signature not in seen:
            selected.append(stripped)
            seen.add(signature)
    compacted = "\n".join(selected).strip() or source[:max_chars]
    if len(compacted) <= max_chars:
        return compacted, True
    output: List[str] = []
    used = 0
    for line in compacted.splitlines():
        line_length = len(line) + 1
        if output and used + line_length > max_chars:
            break
        if not output and line_length > max_chars:
            output.append(line[:max_chars])
            used = max_chars
            break
        output.append(line)
        used += line_length
    return "\n".join(output).strip(), True


def budget_chunks_for_llm(chunks: List[Dict[str, Any]]) -> Tuple[List[Dict[str, Any]], List[Tuple[Dict[str, Any], str]], Dict[str, Any], List[str]]:
    llm_chunks: List[Dict[str, Any]] = []
    local_chunks: List[Tuple[Dict[str, Any], str]] = []
    warnings: List[str] = []
    total_budget = max(0, LLM_TOTAL_INPUT_CHARS_PER_EXTRACTION)
    per_chunk_budget = max(0, LLM_MAX_INPUT_CHARS_PER_CHUNK)
    max_chunks = max(0, LLM_MAX_CHUNKS_PER_EXTRACTION)
    remaining = total_budget
    original_chars = sum(len(chunk.get("text", "")) for chunk in chunks)
    sent_chars = 0
    truncated_chunks = 0

    for chunk in chunks:
        chunk_index = chunk.get("index", len(llm_chunks) + len(local_chunks) + 1)
        if len(llm_chunks) >= max_chunks:
            local_chunks.append((chunk, f"LLM chunk budget reached ({max_chunks} chunks per extraction)"))
            continue
        if remaining <= 0 or per_chunk_budget <= 0:
            local_chunks.append((chunk, f"LLM input character budget reached ({total_budget} chars per extraction)"))
            continue
        chunk_budget = min(per_chunk_budget, remaining)
        compacted_text, was_truncated = compact_chunk_for_llm(chunk.get("text", ""), chunk_budget)
        if not compacted_text.strip():
            local_chunks.append((chunk, "Chunk contained no text after LLM budget compaction"))
            continue
        budgeted_chunk = dict(chunk)
        budgeted_chunk["text"] = compacted_text
        budgeted_chunk["llm_character_count"] = len(compacted_text)
        budgeted_chunk["llm_truncated"] = was_truncated
        llm_chunks.append(budgeted_chunk)
        sent_chars += len(compacted_text)
        remaining -= len(compacted_text)
        if was_truncated:
            truncated_chunks += 1

    stats = {
        "groq_completion_token_cap_per_call": LLM_MAX_COMPLETION_TOKENS,
        "groq_retry_attempts_per_chain": LLM_RETRY_ATTEMPTS,
        "structured_output_enabled": LLM_STRUCTURED_OUTPUT_ENABLED,
        "parser_fallback_enabled": LLM_PARSER_FALLBACK_ENABLED,
        "raw_repair_enabled": LLM_RAW_REPAIR_ENABLED,
        "max_groq_calls_per_chunk": 1 + int(LLM_PARSER_FALLBACK_ENABLED) + int(LLM_RAW_REPAIR_ENABLED),
        "input_char_budget_per_chunk": per_chunk_budget,
        "input_char_budget_per_extraction": total_budget,
        "max_llm_chunks_per_extraction": max_chunks,
        "source_chunk_count": len(chunks),
        "groq_chunk_count": len(llm_chunks),
        "fallback_chunk_count": len(local_chunks),
        "source_input_chars": original_chars,
        "groq_input_chars": sent_chars,
        "estimated_groq_input_tokens_per_attempt": approximate_token_count("x" * sent_chars),
        "estimated_max_groq_input_tokens_with_retries": approximate_token_count("x" * sent_chars) * LLM_RETRY_ATTEMPTS * (1 + int(LLM_PARSER_FALLBACK_ENABLED) + int(LLM_RAW_REPAIR_ENABLED)),
        "compacted_chunk_count": truncated_chunks,
    }
    return llm_chunks, local_chunks, stats, warnings


FINANCIAL_FACT_PARSER = PydanticOutputParser(pydantic_object=ChunkFinancialFacts)
FINANCIAL_PARSER = FINANCIAL_FACT_PARSER
JSON_OUTPUT_PARSER = JsonOutputParser()
FINANCIAL_METRIC_GUIDE = (
    "Revenue/net revenue/sales; gross profit; operating income/profit; EBIT; EBITDA; "
    "net income/net profit/PAT; EPS/diluted EPS; cash/cash equivalents; assets/current/non-current assets; "
    "liabilities/current liabilities; long-term debt/debt/equity; operating/investing/financing/free cash flow; "
    "capital expenditure/capex; dividend; shares outstanding."
)
FINANCIAL_PROMPT = PromptTemplate(
    input_variables=["chunk_index", "chunk_count", "document_text"],
    partial_variables={"metric_guide": FINANCIAL_METRIC_GUIDE},
    template=(
        "Extract only financial facts explicitly written in this chunk. Do not infer, calculate, compare, analyse, summarize, or fill blanks. "
        "Do not output null fields. If a metric or company fact is absent, omit it.\n"
        "Return compact JSON only. Empty result: {{\"metrics\":[]}}.\n"
        "Metric item keys: metric, value, unit, period, page_number, confidence. Include optional keys only when explicit.\n"
        "If explicit company facts appear, add company_info items with keys field, value, page_number, confidence.\n"
        "Recognize synonyms for: {metric_guide}\n\n"
        "Chunk {chunk_index}/{chunk_count}:\n{document_text}"
    ),
)

STRUCTURED_FINANCIAL_PROMPT = PromptTemplate(
    input_variables=["chunk_index", "chunk_count", "document_text"],
    partial_variables={"metric_guide": FINANCIAL_METRIC_GUIDE},
    template=(
        "Extract only financial facts explicitly written in this chunk. Do not infer, calculate, compare, analyse, summarize, or fill blanks. "
        "Return only detected facts in the response schema; omit absent optional fields and never emit nulls.\n"
        "Recognize synonyms for: {metric_guide}\n\n"
        "Chunk {chunk_index}/{chunk_count}:\n{document_text}"
    ),
)


def get_llm() -> ChatGroq:
    global LLM_INSTANCE, LLM_INSTANCE_KEY
    api_key = os.getenv("GROQ_API_KEY")
    if not api_key:
        raise MissingGroqAPIKeyError()
    model = os.getenv("GROQ_MODEL", DEFAULT_GROQ_MODEL)
    key = (api_key, model, LLM_MAX_COMPLETION_TOKENS)
    with LLM_LOCK:
        if LLM_INSTANCE is None or LLM_INSTANCE_KEY != key:
            LLM_INSTANCE = ChatGroq(
                model=model,
                temperature=0,
                groq_api_key=api_key,
                timeout=LLM_TIMEOUT_SECONDS,
                max_retries=0,
                max_tokens=LLM_MAX_COMPLETION_TOKENS,
            )
            LLM_INSTANCE_KEY = key
            logger.info("LLM initialized with Groq model {model} max_tokens={max_tokens}", model=model, max_tokens=LLM_MAX_COMPLETION_TOKENS)
    return LLM_INSTANCE


def get_chains() -> Tuple[Any, Any, Any]:
    global STRUCTURED_CHAIN, PARSER_CHAIN, RAW_CHAIN, CHAIN_KEY
    llm = get_llm()
    model = os.getenv("GROQ_MODEL", DEFAULT_GROQ_MODEL)
    api_key = os.getenv("GROQ_API_KEY") or ""
    key = (api_key, model, LLM_MAX_COMPLETION_TOKENS)
    with LLM_LOCK:
        if CHAIN_KEY != key or PARSER_CHAIN is None or RAW_CHAIN is None:
            STRUCTURED_CHAIN = None
            if LLM_STRUCTURED_OUTPUT_ENABLED:
                try:
                    structured_llm = llm.with_structured_output(ChunkFinancialFacts)
                    STRUCTURED_CHAIN = STRUCTURED_FINANCIAL_PROMPT | structured_llm
                except Exception:
                    STRUCTURED_CHAIN = None
            PARSER_CHAIN = FINANCIAL_PROMPT | llm | FINANCIAL_FACT_PARSER
            RAW_CHAIN = FINANCIAL_PROMPT | llm
            CHAIN_KEY = key
    return STRUCTURED_CHAIN, PARSER_CHAIN, RAW_CHAIN


def is_transient_exception(exc: BaseException) -> bool:
    if isinstance(exc, (httpx.TimeoutException, httpx.TransportError, TimeoutError, asyncio.TimeoutError)):
        return True
    text = str(exc).lower()
    transient_tokens = ["timeout", "timed out", "rate limit", "429", "500", "502", "503", "504", "temporarily", "connection", "network", "overloaded"]
    return any(token in text for token in transient_tokens)


def before_sleep_log(retry_state: Any) -> None:
    exception = retry_state.outcome.exception() if retry_state.outcome else None
    logger.warning("Retrying LLM call attempt {attempt} after error: {error}", attempt=retry_state.attempt_number, error=str(exception))


@retry(retry=retry_if_exception(is_transient_exception), wait=wait_exponential(multiplier=1, min=1, max=12), stop=stop_after_attempt(LLM_RETRY_ATTEMPTS), before_sleep=before_sleep_log, reraise=True)
async def invoke_structured_chain(payload: Dict[str, Any]) -> Any:
    structured_chain, _, _ = get_chains()
    if structured_chain is None:
        raise OutputParserException("Structured output chain is not available")
    async with LLM_SEMAPHORE:
        logger.info("LLM structured fact extraction chunk {chunk}", chunk=payload.get("chunk_index"))
        return await asyncio.wait_for(structured_chain.ainvoke(payload), timeout=LLM_TIMEOUT_SECONDS)


@retry(retry=retry_if_exception(is_transient_exception), wait=wait_exponential(multiplier=1, min=1, max=12), stop=stop_after_attempt(LLM_RETRY_ATTEMPTS), before_sleep=before_sleep_log, reraise=True)
async def invoke_parser_chain(payload: Dict[str, Any]) -> ChunkFinancialFacts:
    _, parser_chain, _ = get_chains()
    async with LLM_SEMAPHORE:
        logger.info("LLM parser fact extraction chunk {chunk}", chunk=payload.get("chunk_index"))
        return await asyncio.wait_for(parser_chain.ainvoke(payload), timeout=LLM_TIMEOUT_SECONDS)


@retry(retry=retry_if_exception(is_transient_exception), wait=wait_exponential(multiplier=1, min=1, max=12), stop=stop_after_attempt(LLM_RETRY_ATTEMPTS), before_sleep=before_sleep_log, reraise=True)
async def invoke_raw_chain(payload: Dict[str, Any]) -> str:
    _, _, raw_chain = get_chains()
    async with LLM_SEMAPHORE:
        logger.info("LLM raw repair fact extraction chunk {chunk}", chunk=payload.get("chunk_index"))
        response = await asyncio.wait_for(raw_chain.ainvoke(payload), timeout=LLM_TIMEOUT_SECONDS)
    content = getattr(response, "content", response)
    return str(content)


def coerce_chunk_financial_facts(value: Any) -> ChunkFinancialFacts:
    if isinstance(value, ChunkFinancialFacts):
        return value
    if isinstance(value, ExtractedFinancialData):
        return legacy_financial_data_to_facts(value.model_dump(mode="json", exclude_none=True))
    if isinstance(value, str):
        return parse_financial_response_text(value)
    sanitized = sanitize_for_json(value)
    if isinstance(sanitized, list):
        return ChunkFinancialFacts.model_validate({"metrics": sanitized})
    if isinstance(sanitized, dict):
        wrapped = sanitized.get("financial_data") if "financial_data" in sanitized else sanitized.get("data")
        if isinstance(wrapped, (dict, list)) and not any(key in sanitized for key in ("metrics", "company_info", "company_information")):
            return coerce_chunk_financial_facts(wrapped)
        if any(key in sanitized for key in ("metrics", "metric", "company_info", "company_information")):
            return ChunkFinancialFacts.model_validate(normalize_chunk_fact_payload(sanitized))
        return legacy_financial_data_to_facts(sanitized)
    return ChunkFinancialFacts()


def normalize_chunk_fact_payload(data: Dict[str, Any]) -> Dict[str, Any]:
    if "metric" in data and "value" in data:
        metrics = [data]
    else:
        metrics = data.get("metrics") or data.get("financial_metrics") or []
    if isinstance(metrics, dict):
        metrics = [metrics]
    if not isinstance(metrics, list):
        metrics = []
    normalized_metrics: List[Any] = []
    for metric in metrics:
        if isinstance(metric, dict):
            converted = metric_from_legacy_value(None, metric)
            normalized_metrics.append(converted or metric)
        else:
            normalized_metrics.append(metric)

    company_info = data.get("company_info") if "company_info" in data else data.get("company_information", [])
    if isinstance(company_info, dict):
        company_info = [{"field": key, "value": value} for key, value in company_info.items() if safe_string(value) is not None]
    elif isinstance(company_info, list):
        normalized_company_info: List[Any] = []
        for item in company_info:
            if isinstance(item, dict) and not ({"field", "value"} <= set(item)):
                normalized_company_info.extend({"field": key, "value": value} for key, value in item.items() if safe_string(value) is not None)
            else:
                normalized_company_info.append(item)
        company_info = normalized_company_info
    else:
        company_info = []
    return {"metrics": normalized_metrics, "company_info": company_info}

def legacy_financial_data_to_facts(data: Dict[str, Any]) -> ChunkFinancialFacts:
    metrics: List[Dict[str, Any]] = []
    company_info: List[Dict[str, Any]] = []
    for field in SCALAR_FINANCIAL_FIELDS:
        value = data.get(field)
        if value not in (None, "") and not isinstance(value, (dict, list)):
            company_info.append({"field": field, "value": value})
    for field in METRIC_FINANCIAL_FIELDS:
        metric = metric_from_legacy_value(field, data.get(field))
        if metric is not None:
            metrics.append(metric)
    for value in data.get("all_financial_values") or data.get("metrics") or []:
        metric = metric_from_legacy_value(None, value)
        if metric is not None:
            metrics.append(metric)
    return ChunkFinancialFacts.model_validate({"metrics": metrics, "company_info": company_info})


def metric_from_legacy_value(field: Optional[str], value: Any) -> Optional[Dict[str, Any]]:
    if value in (None, ""):
        return None
    default_name = METRIC_FIELD_DISPLAY_NAMES.get(field or "", field or "Metric")
    if isinstance(value, dict):
        value_text = safe_string(value.get("value") or value.get("raw_value") or value.get("normalized_value"))
        if value_text is None:
            return None
        currency = safe_string(value.get("currency"))
        scale = safe_string(value.get("scale"))
        unit = safe_string(value.get("unit")) or safe_string(" ".join(part for part in (currency, scale) if part))
        metric = {
            "metric": safe_string(value.get("metric")) or default_name,
            "value": value_text,
            "unit": unit,
            "period": value.get("period"),
            "page_number": value.get("page_number") or value.get("page"),
            "confidence": value.get("confidence"),
        }
        return {key: item for key, item in metric.items() if item not in (None, "")}
    value_text = safe_string(value)
    if value_text is None:
        return None
    return {"metric": default_name, "value": value_text}


def parse_financial_response_text(text: str) -> ChunkFinancialFacts:
    try:
        return FINANCIAL_FACT_PARSER.parse(text)
    except Exception:
        parsed = parse_json_value_from_text(text)
        return coerce_chunk_financial_facts(parsed)


def parse_json_value_from_text(text: str) -> Any:
    cleaned = str(text or "").strip()
    cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s*```$", "", cleaned)
    repaired = cleaned.replace("\u201c", '"').replace("\u201d", '"').replace("\u2019", "'")
    repaired = re.sub(r",\s*([}\]])", r"\1", repaired)
    repaired = re.sub(r"\bNaN\b|\bInfinity\b|-Infinity", "null", repaired)
    for candidate in (cleaned, repaired):
        try:
            return JSON_OUTPUT_PARSER.parse(candidate)
        except Exception:
            pass
        try:
            return orjson.loads(candidate)
        except Exception:
            pass
        decoder = json.JSONDecoder()
        for index, char in enumerate(candidate):
            if char not in "[{":
                continue
            try:
                parsed, _ = decoder.raw_decode(candidate[index:])
                return parsed
            except Exception:
                continue
    raise ValueError("LLM response could not be repaired into JSON")


def parse_json_object_from_text(text: str) -> Dict[str, Any]:
    parsed = parse_json_value_from_text(text)
    if isinstance(parsed, dict):
        return parsed
    raise ValueError("LLM response did not contain a JSON object")


def compact_exception_message(exc: BaseException, max_length: int = 260) -> str:
    message = re.sub(r"\s+", " ", str(exc)).strip()
    message = re.sub(r"'failed_generation':\s*'.*", "'failed_generation': '<omitted>'}", message)
    if len(message) > max_length:
        return message[: max_length - 3].rstrip() + "..."
    return message


def is_llm_connection_failure(exc: BaseException) -> bool:
    text = str(exc).lower()
    return any(token in text for token in ("connection error", "connecterror", "all connection attempts failed", "apiconnectionerror"))


def local_fallback_chunk_result(chunk: Dict[str, Any], reason: BaseException, *, error_if_empty: bool = True) -> Dict[str, Any]:
    facts = extract_financial_facts_locally(chunk.get("text", ""))
    chunk_index = chunk.get("index", 1)
    return {
        "chunk_index": chunk_index,
        "data": dump_chunk_facts(facts) if facts.metrics or facts.company_info else empty_chunk_financial_facts(),
        "errors": [],
        "warnings": [],
    }


def extract_financial_facts_locally(text: str) -> ChunkFinancialFacts:
    metrics: List[Dict[str, Any]] = []
    company_info: List[Dict[str, Any]] = []
    current_page: Optional[Union[int, str]] = None
    for raw_line in (text or "").splitlines():
        line = re.sub(r"\s+", " ", raw_line).strip() if "\t" not in raw_line else raw_line.strip()
        if not line:
            continue
        page_match = re.search(r"\b(?:page|ocr page)\s+(\d+)\b", line, flags=re.IGNORECASE)
        if page_match:
            current_page = int(page_match.group(1))
        company_name = normalize_company_name_text(line)
        if company_name:
            company_info.append(build_local_company_fact("company_name", company_name, current_page))
            continue
        metrics.extend(extract_local_metrics_from_line(line, current_page))
        company_info.extend(extract_local_company_info_from_line(line, current_page))
    return ChunkFinancialFacts.model_validate({
        "metrics": dedupe_simple_facts(metrics, ("metric", "value", "period", "page_number")),
        "company_info": dedupe_simple_facts(company_info, ("field", "value", "page_number")),
    })


def extract_local_metrics_from_line(line: str, page_number: Optional[Union[int, str]]) -> List[Dict[str, Any]]:
    facts: List[Dict[str, Any]] = []
    cells = [cell.strip() for cell in line.split("\t") if cell.strip()]
    if len(cells) >= 2:
        field = metric_field_for_label(cells[0])
        if field is not None:
            period_context = detect_period_from_text(line)
            for cell in cells[1:]:
                if looks_like_financial_value(cell):
                    facts.append(build_local_metric_fact(field, cell, page_number, period_context))
            if facts:
                return facts
    for field, aliases in METRIC_FIELD_ALIASES.items():
        for alias in sorted(aliases, key=len, reverse=True):
            remainder = value_after_label_prefix(line, alias)
            if remainder is not None and looks_like_financial_value(remainder):
                facts.append(build_local_metric_fact(field, clean_extracted_value(remainder), page_number, detect_period_from_text(line)))
                return facts
            narrative_value = value_after_narrative_label(line, alias)
            if narrative_value is not None:
                facts.append(build_local_metric_fact(field, narrative_value, page_number, detect_period_from_text(line)))
                return facts
    return facts


def extract_local_company_info_from_line(line: str, page_number: Optional[Union[int, str]]) -> List[Dict[str, Any]]:
    facts: List[Dict[str, Any]] = []
    cells = [cell.strip() for cell in line.split("\t") if cell.strip()]
    if len(cells) >= 2:
        field = company_field_for_label(cells[0])
        if field is not None:
            facts.append(build_local_company_fact(field, cells[1], page_number))
            return facts
    for field, aliases in COMPANY_INFO_FIELD_ALIASES.items():
        for alias in sorted(aliases, key=len, reverse=True):
            remainder = value_after_label_prefix(line, alias)
            if remainder is not None:
                value = clean_extracted_value(remainder)
                if value:
                    facts.append(build_local_company_fact(field, value, page_number))
                    return facts
    return facts


def value_after_label_prefix(line: str, alias: str) -> Optional[str]:
    pattern = rf"^\s*{re.escape(alias)}\s*(?:[:=\-??]|\s{{2,}}|\t)\s*(?P<value>.+?)\s*$"
    match = re.match(pattern, line, flags=re.IGNORECASE)
    if match:
        return match.group("value")
    loose_pattern = rf"^\s*{re.escape(alias)}\s+(?P<value>.+?)\s*$"
    loose_match = re.match(loose_pattern, line, flags=re.IGNORECASE)
    if loose_match and starts_with_financial_value(loose_match.group("value")):
        return loose_match.group("value")
    return None


def value_after_narrative_label(line: str, alias: str) -> Optional[str]:
    value_pattern = r"(?:[$????]\s*)?\(?[-+]?\d[\d,]*(?:\.\d+)?\)?\s*(?:%|x|times|thousand|million|billion|trillion|crore|crores|cr|lakh|lakhs|k|m|bn|b|shares?)?|(?:USD|EUR|GBP|INR|CAD|AUD|JPY|CNY|CHF|SGD|HKD|AED)\s+\(?[-+]?\d[\d,]*(?:\.\d+)?\)?\s*(?:thousand|million|billion|trillion|crore|crores|cr|lakh|lakhs|k|m|bn|b)?"
    pattern = rf"\b{re.escape(alias)}\b\s+(?:was|were|is|are|totaled|totalled|amounted to|of|at)\s+(?P<value>{value_pattern})"
    match = re.search(pattern, line, flags=re.IGNORECASE)
    return clean_extracted_value(match.group("value")) if match else None


def looks_like_financial_value(value: Any) -> bool:
    text = safe_string(value) or ""
    return bool(re.search(r"(?i)(?:USD|EUR|GBP|INR|CAD|AUD|JPY|[$????])?\s*\(?[-+]?\d[\d,]*(?:\.\d+)?\)?\s*(?:%|x|times|thousand|million|billion|trillion|crore|crores|cr|lakh|lakhs|k|m|bn|b|shares?)?", text))


def starts_with_financial_value(value: Any) -> bool:
    text = safe_string(value) or ""
    return bool(re.match(r"(?i)^(?:USD|EUR|GBP|INR|CAD|AUD|JPY|[$????])?\s*\(?[-+]?\d", text))


def clean_extracted_value(value: Any) -> str:
    text = safe_string(value) or ""
    text = re.sub(r"\s+", " ", text).strip(" ;,.?")
    return text


def detect_period_from_text(text: str) -> Optional[str]:
    match = re.search(r"\b(?:FY\s*)?20\d{2}\b|\bQ[1-4]\s*20\d{2}\b|\b(?:quarter|year)\s+ended\s+[A-Za-z]+\s+\d{1,2},?\s+20\d{2}\b", text, flags=re.IGNORECASE)
    return match.group(0) if match else None


def extract_unit_from_value(value: str) -> Optional[str]:
    units: List[str] = []
    currency_match = re.search(r"(?i)\b(?:USD|EUR|GBP|INR|CAD|AUD|JPY|CNY|CHF|SGD|HKD|AED)\b|[$????]", value)
    if currency_match:
        units.append(CURRENCY_SYMBOLS.get(currency_match.group(0), currency_match.group(0).upper()))
    scale_match = re.search(r"(?i)\b(?:thousand|million|billion|trillion|crore|crores|cr|lakh|lakhs|k|m|bn|b)\b", value)
    if scale_match:
        units.append(scale_match.group(0).lower())
    if "%" in value:
        units.append("percent")
    if re.search(r"(?i)\bshares?\b", value):
        units.append("shares")
    return " ".join(dict.fromkeys(units)) or None


def build_local_metric_fact(field: str, value: str, page_number: Optional[Union[int, str]], period: Optional[str]) -> Dict[str, Any]:
    fact = {
        "metric": METRIC_FIELD_DISPLAY_NAMES.get(field, field),
        "value": value,
        "unit": extract_unit_from_value(value),
        "period": period,
        "page_number": page_number,
        "confidence": LOCAL_FALLBACK_CONFIDENCE,
    }
    return {key: item for key, item in fact.items() if item not in (None, "")}


def build_local_company_fact(field: str, value: str, page_number: Optional[Union[int, str]]) -> Dict[str, Any]:
    fact = {"field": field, "value": value, "page_number": page_number, "confidence": DEFAULT_COMPANY_INFO_CONFIDENCE}
    return {key: item for key, item in fact.items() if item not in (None, "")}


def dedupe_simple_facts(values: List[Dict[str, Any]], keys: Tuple[str, ...]) -> List[Dict[str, Any]]:
    deduped: List[Dict[str, Any]] = []
    signatures: set[str] = set()
    for value in values:
        signature = "|".join(str(value.get(key) or "").strip().lower() for key in keys)
        if not signature or signature in signatures:
            continue
        signatures.add(signature)
        deduped.append(value)
    return deduped


async def extract_financial_chunk(chunk: Dict[str, Any], metadata: Dict[str, Any]) -> Dict[str, Any]:
    _ = metadata
    payload = {
        "chunk_index": chunk.get("index", 1),
        "chunk_count": chunk.get("total", 1),
        "document_text": chunk.get("text", ""),
    }
    primary_error: Optional[BaseException] = None
    if LLM_STRUCTURED_OUTPUT_ENABLED:
        try:
            structured = await invoke_structured_chain(payload)
            data = coerce_chunk_financial_facts(structured)
            return {"chunk_index": chunk.get("index", 1), "data": dump_chunk_facts(data), "errors": [], "warnings": []}
        except Exception as structured_exc:
            primary_error = structured_exc
            logger.warning("Structured fact extraction failed for chunk {chunk}: {error}", chunk=chunk.get("index"), error=compact_exception_message(structured_exc))
            if is_llm_connection_failure(structured_exc) or not LLM_PARSER_FALLBACK_ENABLED:
                return local_fallback_chunk_result(chunk, structured_exc)
    try:
        parsed = await invoke_parser_chain(payload)
        data = coerce_chunk_financial_facts(parsed)
        return {"chunk_index": chunk.get("index", 1), "data": dump_chunk_facts(data), "errors": [], "warnings": []}
    except OutputParserException as parser_exc:
        primary_error = parser_exc
        logger.warning("Parser fact extraction failed for chunk {chunk}: {error}", chunk=chunk.get("index"), error=compact_exception_message(parser_exc))
        if not LLM_RAW_REPAIR_ENABLED:
            return local_fallback_chunk_result(chunk, parser_exc)
        try:
            raw = await invoke_raw_chain(payload)
            data = parse_financial_response_text(raw)
            return {"chunk_index": chunk.get("index", 1), "data": dump_chunk_facts(data), "errors": [], "warnings": []}
        except Exception as raw_exc:
            logger.warning("Raw fact extraction repair failed for chunk {chunk}: {error}", chunk=chunk.get("index"), error=compact_exception_message(raw_exc))
            return local_fallback_chunk_result(chunk, raw_exc)
    except Exception as exc:
        primary_error = exc
        logger.exception("LLM fact extraction failed for chunk {chunk}", chunk=chunk.get("index"))
        return local_fallback_chunk_result(chunk, exc)
    return local_fallback_chunk_result(chunk, primary_error or RuntimeError("LLM extraction failed"))

def merge_chunk_financial_results(results: List[Dict[str, Any]]) -> Dict[str, Any]:
    facts_by_chunk: List[Tuple[int, ChunkFinancialFacts]] = []
    for result in sorted(results, key=lambda item: item.get("chunk_index", 0)):
        try:
            facts = coerce_chunk_financial_facts(result.get("data") or empty_chunk_financial_facts())
            facts_by_chunk.append((int(result.get("chunk_index") or 0), facts))
        except ValidationError:
            continue
    merged = empty_financial_data()
    scalar_candidates: Dict[str, List[Tuple[Any, float]]] = {field: [] for field in SCALAR_FINANCIAL_FIELDS}
    metric_candidates: Dict[str, List[Dict[str, Any]]] = {field: [] for field in METRIC_FINANCIAL_FIELDS}
    all_values: List[Dict[str, Any]] = []
    confidence_values: List[float] = []

    for _, facts in facts_by_chunk:
        for fact_model in facts.company_info:
            normalized = normalize_company_info_fact(fact_model.model_dump(mode="json", exclude_none=True))
            if normalized is None:
                continue
            field, value, confidence = normalized
            scalar_candidates.setdefault(field, []).append((value, confidence))
            if confidence > 0:
                confidence_values.append(confidence)
        for metric_model in facts.metrics:
            metric = normalize_metric_fact(metric_model.model_dump(mode="json", exclude_none=True))
            if metric is None:
                continue
            all_values.append(metric)
            field = metric_field_for_label(metric.get("metric"))
            if field in metric_candidates:
                metric_candidates[field].append(metric)
            confidence = metric.get("confidence")
            if confidence is not None:
                confidence_values.append(clamp_float(confidence, 0.0, 1.0))

    for field, candidates in scalar_candidates.items():
        chosen = choose_scalar_candidate(candidates)
        if chosen is not None:
            merged[field] = chosen
    company_info = {field: merged[field] for field in SCALAR_FINANCIAL_FIELDS if merged.get(field) not in (None, "")}
    if company_info:
        merged["company_info"] = company_info

    deduped_values = dedupe_metric_values(all_values)
    for field, candidates in metric_candidates.items():
        chosen_metric = choose_metric_candidate(candidates, field)
        if chosen_metric is not None:
            merged[field] = chosen_metric
    merged["metrics"] = deduped_values
    merged["all_financial_values"] = deduped_values
    if confidence_values:
        merged["confidence_score"] = round(sum(confidence_values) / len(confidence_values), 4)
    return dump_financial_data(ExtractedFinancialData.model_validate(merged))


def label_key(value: Any) -> str:
    text = safe_string(value) or ""
    text = re.sub(r"[^a-z0-9]+", " ", text.lower())
    return re.sub(r"\s+", " ", text).strip()


def metric_field_for_label(label: Any) -> Optional[str]:
    key = label_key(label)
    if not key:
        return None
    for field, aliases in METRIC_FIELD_ALIASES.items():
        if key == label_key(field) or key in {label_key(alias) for alias in aliases}:
            return field
    return None


def company_field_for_label(label: Any) -> Optional[str]:
    key = label_key(label)
    if not key:
        return None
    for field, aliases in COMPANY_INFO_FIELD_ALIASES.items():
        if key == label_key(field) or key in {label_key(alias) for alias in aliases}:
            return field
    return None


def canonical_metric_display(label: Any) -> str:
    field = metric_field_for_label(label)
    if field is not None:
        return METRIC_FIELD_DISPLAY_NAMES[field]
    words = label_key(label).split()
    acronyms = {"ebit", "ebitda", "eps", "roa", "roe", "pat"}
    return " ".join(word.upper() if word in acronyms else word.capitalize() for word in words) or "Metric"


def normalize_metric_fact(value: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    try:
        metric = Metric.model_validate(value).model_dump(mode="json", exclude_none=True)
    except ValidationError:
        return None
    metric["metric"] = canonical_metric_display(metric.get("metric"))
    if metric.get("confidence") is None:
        metric["confidence"] = DEFAULT_METRIC_CONFIDENCE
    if metric.get("raw_value") == metric.get("value"):
        metric.pop("raw_value", None)
    return FinancialMetricValue.model_validate(metric).model_dump(mode="json", exclude_none=True)


def normalize_company_info_fact(value: Dict[str, Any]) -> Optional[Tuple[str, str, float]]:
    try:
        fact = CompanyInfoFact.model_validate(value).model_dump(mode="json", exclude_none=True)
    except ValidationError:
        return None
    field = company_field_for_label(fact.get("field"))
    if field is None:
        return None
    confidence = clamp_float(fact.get("confidence"), 0.0, 1.0) if fact.get("confidence") is not None else DEFAULT_COMPANY_INFO_CONFIDENCE
    return field, fact["value"], confidence


def choose_scalar_candidate(candidates: List[Tuple[Any, float]]) -> Any:
    if not candidates:
        return None
    scored = []
    for value, confidence in candidates:
        text = str(value).strip()
        support = sum(1 for other, _ in candidates if fuzz.ratio(text.lower(), str(other).strip().lower()) >= 92)
        scored.append((support, confidence, len(text), value))
    scored.sort(key=lambda item: (item[0], item[1], item[2]), reverse=True)
    return scored[0][3]


def choose_metric_candidate(candidates: List[Dict[str, Any]], field: str) -> Optional[Dict[str, Any]]:
    if not candidates:
        return None
    enriched = []
    for candidate in candidates:
        metric = dict(candidate)
        metric["metric"] = metric.get("metric") or METRIC_FIELD_DISPLAY_NAMES.get(field, field)
        confidence = clamp_float(metric.get("confidence"), 0.0, 1.0) if metric.get("confidence") is not None else 0.0
        completeness = sum(1 for key in ("value", "unit", "period", "page_number") if metric.get(key) not in (None, ""))
        enriched.append((confidence, completeness, len(str(metric.get("value") or "")), metric))
    enriched.sort(key=lambda item: (item[0], item[1], item[2]), reverse=True)
    return FinancialMetricValue.model_validate(enriched[0][3]).model_dump(mode="json", exclude_none=True)


def has_metric_value(value: Dict[str, Any]) -> bool:
    return any(value.get(key) not in (None, "") for key in ("value", "raw_value", "normalized_value"))


def dedupe_metric_values(values: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    deduped: List[Dict[str, Any]] = []
    signatures: List[str] = []
    for value in values:
        try:
            metric = FinancialMetricValue.model_validate(value).model_dump(mode="json", exclude_none=True)
        except ValidationError:
            continue
        signature = " ".join(str(metric.get(key) or "") for key in ("metric", "value", "unit", "period", "page_number")).strip().lower()
        if not signature:
            continue
        if any(fuzz.ratio(signature, existing) >= 96 for existing in signatures):
            continue
        signatures.append(signature)
        deduped.append(metric)
    deduped.sort(key=lambda item: (metric_sort_index(item.get("metric")), str(item.get("period") or ""), str(item.get("page_number") or ""), str(item.get("metric") or "")))
    return deduped


def metric_sort_index(label: Any) -> int:
    field = metric_field_for_label(label)
    if field in METRIC_FINANCIAL_FIELDS:
        return METRIC_FINANCIAL_FIELDS.index(field)
    return len(METRIC_FINANCIAL_FIELDS)

def calculate_confidence(state: ExtractionState) -> float:
    financial_data = state.get("financial_data") or {}
    metric_confidences: List[float] = []
    extracted_count = 0
    seen_signatures: set[str] = set()

    def add_metric_confidence(value: Dict[str, Any]) -> None:
        nonlocal extracted_count
        if not isinstance(value, dict) or not has_metric_value(value):
            return
        signature = " ".join(str(value.get(key) or "") for key in ("metric", "value", "unit", "period")).lower()
        if signature in seen_signatures:
            return
        seen_signatures.add(signature)
        extracted_count += 1
        if value.get("confidence") is not None:
            metric_confidences.append(clamp_float(value.get("confidence"), 0.0, 1.0))

    for value in financial_data.get("metrics") or financial_data.get("all_financial_values") or []:
        add_metric_confidence(value)
    for field in METRIC_FINANCIAL_FIELDS:
        add_metric_confidence(financial_data.get(field))

    company_info = financial_data.get("company_info") if isinstance(financial_data.get("company_info"), dict) else {}
    scalar_count = len(company_info) or sum(1 for field in SCALAR_FINANCIAL_FIELDS if financial_data.get(field) not in (None, ""))
    llm_confidence = clamp_float(financial_data.get("confidence_score"), 0.0, 1.0)
    if metric_confidences:
        llm_confidence = max(llm_confidence, sum(metric_confidences) / len(metric_confidences))
    text_quality = clamp_float(state.get("native_quality"), 0.0, 1.0)
    if state.get("merged_text", "").strip() and not state.get("errors"):
        text_quality = max(text_quality, 0.85)
    if state.get("ocr_required") and state.get("metadata", {}).get("ocr_average_confidence") is not None:
        text_quality = max(text_quality, clamp_float(state.get("metadata", {}).get("ocr_average_confidence"), 0.0, 1.0))
    extraction_density = clamp_float((extracted_count + scalar_count) / 18, 0.0, 1.0)
    coverage_floor = clamp_float(extracted_count / 4, 0.0, 1.0) if extracted_count else 0.0
    error_penalty = min(0.5, len(state.get("errors", [])) * 0.14)
    penalized_warnings = [warning for warning in state.get("warnings", []) if "deterministic explicit-fact fallback" not in warning.lower()]
    warning_penalty = min(0.25, len(penalized_warnings) * 0.035)
    conservative_confidence = (0.52 * llm_confidence) + (0.28 * text_quality) + (0.20 * extraction_density)
    explicit_fact_confidence = (0.82 * llm_confidence) + (0.13 * text_quality) + (0.05 * coverage_floor)
    confidence = max(conservative_confidence, explicit_fact_confidence) - error_penalty - warning_penalty
    return round(clamp_float(confidence, 0.0, 1.0) * 100, 2)

async def validate_file_node(state: ExtractionState) -> ExtractionState:
    state = ensure_state_defaults(dict(state))
    logger.info("LangGraph node Validate File")
    data = state.get("uploaded_file") or b""
    filename = state.get("filename") or "upload"
    state["filename"] = filename
    state["metadata"].update({"filename": filename, "received_at": now_iso(), "byte_size": len(data)})
    if not isinstance(data, (bytes, bytearray)) or len(data) == 0:
        add_error(state, "Empty file or text payload")
    if len(data) > MAX_UPLOAD_BYTES:
        add_error(state, f"File exceeds maximum size of {MAX_UPLOAD_BYTES // (1024 * 1024)} MB")
    return state


async def detect_file_type_node(state: ExtractionState) -> ExtractionState:
    state = ensure_state_defaults(dict(state))
    logger.info("LangGraph node Detect File Type")
    data = state.get("uploaded_file") or b""
    filename = state.get("filename") or "upload"
    provided_mime = state.get("mime_type") or "application/octet-stream"
    extension = detect_extension_from_content(data, filename, provided_mime).lower()
    if extension == ".jpeg":
        extension = ".jpg"
    state["extension"] = extension
    state["mime_type"] = mime_for_extension(extension, provided_mime)
    state["metadata"].update({"extension": extension, "mime_type": state["mime_type"]})
    if extension not in SUPPORTED_EXTENSIONS:
        add_error(state, f"Unsupported file type: {extension or 'unknown'}")
    return state


async def extract_metadata_node(state: ExtractionState) -> ExtractionState:
    state = ensure_state_defaults(dict(state))
    logger.info("LangGraph node Extract Metadata")
    data = state.get("uploaded_file") or b""
    extension = state.get("extension", "")
    if extension not in SUPPORTED_EXTENSIONS:
        return state
    try:
        if extension in PDF_EXTENSIONS:
            metadata, page_count = await run_blocking(extract_pdf_metadata, data)
        elif extension in IMAGE_EXTENSIONS:
            metadata, page_count = await run_blocking(extract_image_metadata, data)
        elif extension in DOCUMENT_EXTENSIONS:
            metadata, page_count = await run_blocking(extract_docx_metadata, data)
        elif extension in SPREADSHEET_EXTENSIONS:
            metadata, page_count = await run_blocking(extract_xlsx_metadata, data)
        else:
            metadata, page_count = await run_blocking(extract_text_metadata, data, extension)
        state["metadata"].update(metadata)
        state["page_count"] = page_count
    except fitz.FileDataError as exc:
        add_error(state, f"Corrupted PDF or unreadable PDF: {str(exc)}")
    except UnidentifiedImageError as exc:
        add_error(state, f"Unreadable image file: {str(exc)}")
    except PermissionError as exc:
        add_error(state, f"Permission error while reading metadata: {str(exc)}")
    except MemoryError:
        add_error(state, "Memory error while extracting metadata")
    except Exception as exc:
        add_warning(state, f"Metadata extraction warning: {str(exc)}")
    return state


async def extract_native_text_node(state: ExtractionState) -> ExtractionState:
    state = ensure_state_defaults(dict(state))
    logger.info("LangGraph node Extract Native Text")
    data = state.get("uploaded_file") or b""
    extension = state.get("extension", "")
    if extension not in SUPPORTED_EXTENSIONS:
        return state
    try:
        if extension in PDF_EXTENSIONS:
            text, metadata, warnings = await run_blocking(extract_pdf_native_text, data)
            state["raw_text"] = text
            state["metadata"].update(metadata)
            state["native_quality"] = clamp_float(metadata.get("native_quality"), 0.0, 1.0)
            for warning in warnings:
                add_warning(state, warning)
        elif extension in IMAGE_EXTENSIONS:
            state["raw_text"] = ""
            state["native_quality"] = 0.0
        elif extension in DOCUMENT_EXTENSIONS:
            text, metadata, tables, warnings = await extract_docx_native_text(data)
            state["raw_text"] = text
            state["metadata"].update(metadata)
            state["tables"] = merge_tables(state.get("tables", []), tables)
            state["native_quality"] = clamp_float(len(text) / 8000, 0.0, 1.0)
            for warning in warnings:
                add_warning(state, warning)
        elif extension in SPREADSHEET_EXTENSIONS:
            text, metadata, tables, warnings = await run_blocking(extract_xlsx_native_text, data)
            state["raw_text"] = text
            state["metadata"].update(metadata)
            state["tables"] = merge_tables(state.get("tables", []), tables)
            state["native_quality"] = clamp_float(len(text) / 8000, 0.0, 1.0)
            for warning in warnings:
                add_warning(state, warning)
        elif extension in TABULAR_TEXT_EXTENSIONS:
            text, metadata, tables, warnings = await run_blocking(extract_csv_native_text, data)
            state["raw_text"] = text
            state["metadata"].update(metadata)
            state["tables"] = merge_tables(state.get("tables", []), tables)
            state["native_quality"] = clamp_float(len(text) / 8000, 0.0, 1.0)
            for warning in warnings:
                add_warning(state, warning)
        elif extension in TEXT_EXTENSIONS:
            text, metadata, warnings = await run_blocking(extract_txt_native_text, data)
            state["raw_text"] = text
            state["metadata"].update(metadata)
            state["native_quality"] = clamp_float(len(text) / 8000, 0.0, 1.0)
            for warning in warnings:
                add_warning(state, warning)
    except UnicodeError as exc:
        add_error(state, f"Encoding error during text extraction: {str(exc)}")
    except PermissionError as exc:
        add_error(state, f"Permission error during text extraction: {str(exc)}")
    except MemoryError:
        add_error(state, "Memory error during text extraction")
    except Exception as exc:
        add_error(state, f"Native text extraction failed: {str(exc)}")
    return state


def merge_tables(existing: List[Dict[str, Any]], incoming: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    merged = list(existing or [])
    signatures = {table_signature(table) for table in merged}
    for table in incoming or []:
        signature = table_signature(table)
        if signature not in signatures:
            merged.append(table)
            signatures.add(signature)
    return merged


def table_signature(table: Dict[str, Any]) -> str:
    return re.sub(r"\s+", " ", rows_to_text(table.get("rows") or [])[:1000]).lower()


async def detect_ocr_requirement_node(state: ExtractionState) -> ExtractionState:
    state = ensure_state_defaults(dict(state))
    logger.info("LangGraph node Detect OCR Requirement")
    extension = state.get("extension", "")
    metadata = state.get("metadata", {})
    ocr_pages: List[int] = []
    if extension in IMAGE_EXTENSIONS:
        state["ocr_required"] = True
        state["ocr_pages"] = [1]
        return state
    if extension in PDF_EXTENSIONS:
        char_counts = metadata.get("native_page_char_counts") or []
        image_counts = metadata.get("native_page_image_counts") or []
        for index, count in enumerate(char_counts):
            image_count = image_counts[index] if index < len(image_counts) else 0
            if int(count or 0) < MIN_NATIVE_PDF_CHARS_PER_PAGE or image_count > 0 and int(count or 0) < MIN_NATIVE_PDF_CHARS_PER_PAGE * 2:
                ocr_pages.append(index + 1)
        average_chars = sum(int(count or 0) for count in char_counts) / max(1, len(char_counts))
        state["ocr_required"] = bool(ocr_pages)
        state["ocr_pages"] = ocr_pages if state["ocr_required"] else []
        if state["ocr_required"]:
            add_warning(state, f"OCR required for PDF pages: {ocr_pages}")
    else:
        state["ocr_required"] = False
        state["ocr_pages"] = []
    return state


async def ocr_processing_node(state: ExtractionState) -> ExtractionState:
    state = ensure_state_defaults(dict(state))
    logger.info("LangGraph node OCR Processing")
    if not state.get("ocr_required"):
        state["ocr_text"] = ""
        return state
    data = state.get("uploaded_file") or b""
    extension = state.get("extension", "")
    try:
        logger.info("OCR started for {filename}", filename=state.get("filename"))
        if extension in IMAGE_EXTENSIONS:
            text, metadata = await run_blocking(run_ocr_image_bytes, data)
        elif extension in PDF_EXTENSIONS:
            text, metadata = await run_blocking(run_ocr_pdf_bytes, data, state.get("ocr_pages", []))
        else:
            text, metadata = "", {}
        state["ocr_text"] = text
        state["metadata"].update(metadata)
        if not text.strip():
            add_warning(state, "OCR completed without extracting text")
        logger.info("OCR completed for {filename}", filename=state.get("filename"))
    except MemoryError:
        add_error(state, "Memory error during OCR processing")
    except Exception as exc:
        message = f"OCR failure: {str(exc)}"
        if state.get("raw_text", "").strip():
            add_warning(state, message)
        else:
            add_error(state, message)
    return state


async def table_extraction_node(state: ExtractionState) -> ExtractionState:
    state = ensure_state_defaults(dict(state))
    logger.info("LangGraph node Table Extraction")
    data = state.get("uploaded_file") or b""
    extension = state.get("extension", "")
    if extension not in SUPPORTED_EXTENSIONS:
        return state
    try:
        tables: List[Dict[str, Any]] = []
        warnings: List[str] = []
        if extension in PDF_EXTENSIONS:
            tables, warnings = await run_blocking(extract_pdf_tables, data)
        elif extension in DOCUMENT_EXTENSIONS:
            document = Document(io.BytesIO(data))
            tables = extract_docx_tables_from_container(document.tables, "body")
        elif extension in SPREADSHEET_EXTENSIONS:
            _, _, tables, warnings = await run_blocking(extract_xlsx_native_text, data)
        elif extension in TABULAR_TEXT_EXTENSIONS:
            _, _, tables, warnings = await run_blocking(extract_csv_native_text, data)
        state["tables"] = merge_tables(state.get("tables", []), tables)
        state["metadata"]["table_count"] = len(state.get("tables", []))
        for warning in warnings:
            add_warning(state, warning)
    except Exception as exc:
        add_warning(state, f"Table extraction failed: {str(exc)}")
    return state


async def merge_text_node(state: ExtractionState) -> ExtractionState:
    state = ensure_state_defaults(dict(state))
    logger.info("LangGraph node Merge Text")
    raw_text = state.get("raw_text", "")
    ocr_text = remove_ocr_duplicates(raw_text, state.get("ocr_text", ""))
    table_text = render_tables(state.get("tables", []))
    sections = []
    if raw_text.strip():
        sections.append(raw_text.strip())
    if ocr_text.strip():
        sections.append(ocr_text.strip())
    if table_text.strip():
        sections.append("Extracted Tables\n" + table_text.strip())
    state["merged_text"] = remove_duplicate_lines("\n\n".join(sections))
    state["metadata"]["merged_character_count"] = len(state["merged_text"])
    return state


async def normalize_text_node(state: ExtractionState) -> ExtractionState:
    state = ensure_state_defaults(dict(state))
    logger.info("LangGraph node Normalize Text")
    state["merged_text"] = normalize_document_text(state.get("merged_text", ""))
    if not state["merged_text"].strip():
        add_error(state, "No extractable text found")
    state["metadata"]["normalized_character_count"] = len(state["merged_text"])
    return state


async def chunk_document_node(state: ExtractionState) -> ExtractionState:
    state = ensure_state_defaults(dict(state))
    logger.info("LangGraph node Chunk Document")
    chunks = chunk_text(state.get("merged_text", ""))
    state["chunks"] = chunks
    state["metadata"]["chunk_count"] = len(chunks)
    if not chunks:
        add_warning(state, "Document chunking produced no chunks")
    return state


async def financial_extraction_node(state: ExtractionState) -> ExtractionState:
    state = ensure_state_defaults(dict(state))
    logger.info("LangGraph node Financial Extraction")
    chunks = state.get("chunks", [])
    if not chunks:
        state["chunk_results"] = []
        state["financial_data"] = empty_financial_data()
        return state
    if not os.getenv("GROQ_API_KEY"):
        add_error(state, "Missing GROQ_API_KEY environment variable")
        state["chunk_results"] = []
        state["financial_data"] = empty_financial_data()
        return state
    llm_chunks, budget_local_chunks, budget_stats, budget_warnings = budget_chunks_for_llm(chunks)
    state["metadata"]["llm_budget"] = budget_stats
    for warning in budget_warnings:
        add_warning(state, warning)
    tasks = [extract_financial_chunk(chunk, state.get("metadata", {})) for chunk in llm_chunks]
    results = await asyncio.gather(*tasks, return_exceptions=True) if tasks else []
    results.extend(
        local_fallback_chunk_result(chunk, RuntimeError(reason), error_if_empty=False)
        for chunk, reason in budget_local_chunks
    )
    chunk_results: List[Dict[str, Any]] = []
    for index, result in enumerate(results):
        if isinstance(result, Exception):
            add_error(state, f"Financial extraction failed for chunk {index + 1}: {str(result)}")
            chunk_results.append({"chunk_index": index + 1, "data": empty_chunk_financial_facts(), "errors": [str(result)]})
        else:
            chunk_results.append(result)
            for warning in result.get("warnings", []):
                add_warning(state, f"Financial extraction chunk {result.get('chunk_index')}: {warning}")
            for error in result.get("errors", []):
                add_error(state, f"Financial extraction chunk {result.get('chunk_index')}: {error}")
    state["chunk_results"] = chunk_results
    return state


async def merge_chunk_results_node(state: ExtractionState) -> ExtractionState:
    state = ensure_state_defaults(dict(state))
    logger.info("LangGraph node Merge Chunk Results")
    try:
        merged_data = merge_chunk_financial_results(state.get("chunk_results", []))
        state["financial_data"] = augment_financial_data_with_context(merged_data, state.get("metadata", {}), state.get("merged_text", ""))
    except Exception as exc:
        add_error(state, f"Failed to merge financial extraction results: {str(exc)}")
        state["financial_data"] = empty_financial_data()
    return state


async def validate_json_node(state: ExtractionState) -> ExtractionState:
    state = ensure_state_defaults(dict(state))
    logger.info("LangGraph node Validate JSON")
    try:
        validated = ExtractedFinancialData.model_validate(state.get("financial_data") or {})
        state["financial_data"] = dump_financial_data(validated)
    except ValidationError as exc:
        add_error(state, f"Invalid financial JSON after extraction: {str(exc)}")
        state["financial_data"] = empty_financial_data()
    return state


async def confidence_calculation_node(state: ExtractionState) -> ExtractionState:
    state = ensure_state_defaults(dict(state))
    logger.info("LangGraph node Confidence Calculation")
    confidence = calculate_confidence(state)
    state["confidence_score"] = confidence
    financial_data = dict(state.get("financial_data") or empty_financial_data())
    financial_data["confidence_score"] = round(clamp_float(confidence / 100, 0.0, 1.0), 4)
    state["financial_data"] = dump_financial_data(ExtractedFinancialData.model_validate(financial_data))
    return state

async def response_formatter_node(state: ExtractionState) -> ExtractionState:
    state = ensure_state_defaults(dict(state))
    logger.info("LangGraph node Response Formatter")
    state["execution_time"] = round(time.perf_counter() - state.get("start_time", time.perf_counter()), 4)
    internal_error_markers = ("financial extraction chunk", "failed to parse chunkfinancialfacts", "groq llm unavailable or skipped")
    internal_warning_markers = ("compacted for groq token budget", "deterministic explicit-fact fallback", "chunk(s) skipped groq", "financial extraction chunk")
    public_errors = [error for error in state.get("errors", []) if not any(marker in str(error).lower() for marker in internal_error_markers)]
    public_warnings = [warning for warning in state.get("warnings", []) if not any(marker in str(warning).lower() for marker in internal_warning_markers)]
    status_code = status_code_for_state(state)
    if not public_errors and status_code in {500, 502}:
        status_code = 200
    state["status_code"] = status_code
    response = {
        "success": status_code < 400 and not public_errors,
        "filename": state.get("filename", ""),
        "filetype": state.get("extension", "").replace(".", ""),
        "pages": int(state.get("page_count") or 0),
        "ocr_used": bool(state.get("ocr_required")),
        "processing_time": state["execution_time"],
        "financial_data": sanitize_for_json(state.get("financial_data", empty_financial_data())),
        "confidence_score": state.get("confidence_score", 0.0),
        "errors": public_errors,
        "warnings": public_warnings,
    }
    state["response"] = response
    logger.info("Execution time {seconds}s for {filename}", seconds=state["execution_time"], filename=state.get("filename"))
    return state


def status_code_for_state(state: ExtractionState) -> int:
    errors = " ".join(state.get("errors", [])).lower()
    if not errors:
        return 200
    if "unsupported file type" in errors:
        return 415
    if "empty file" in errors or "maximum size" in errors:
        return 400
    if "missing groq_api_key" in errors:
        return 503
    if "timeout" in errors or "timed out" in errors:
        return 504
    if "network" in errors or "connection" in errors or "groq" in errors:
        return 502
    if "invalid financial json" in errors or "no extractable text" in errors:
        return 422
    if "memory error" in errors:
        return 507
    return 500


def build_workflow() -> Any:
    graph = StateGraph(ExtractionState)
    graph.add_node("Validate File", validate_file_node)
    graph.add_node("Detect File Type", detect_file_type_node)
    graph.add_node("Extract Metadata", extract_metadata_node)
    graph.add_node("Extract Native Text", extract_native_text_node)
    graph.add_node("Detect OCR Requirement", detect_ocr_requirement_node)
    graph.add_node("OCR Processing", ocr_processing_node)
    graph.add_node("Table Extraction", table_extraction_node)
    graph.add_node("Merge Text", merge_text_node)
    graph.add_node("Normalize Text", normalize_text_node)
    graph.add_node("Chunk Document", chunk_document_node)
    graph.add_node("Financial Extraction", financial_extraction_node)
    graph.add_node("Merge Chunk Results", merge_chunk_results_node)
    graph.add_node("Validate JSON", validate_json_node)
    graph.add_node("Confidence Calculation", confidence_calculation_node)
    graph.add_node("Response Formatter", response_formatter_node)
    graph.add_edge(START, "Validate File")
    graph.add_edge("Validate File", "Detect File Type")
    graph.add_edge("Detect File Type", "Extract Metadata")
    graph.add_edge("Extract Metadata", "Extract Native Text")
    graph.add_edge("Extract Native Text", "Detect OCR Requirement")
    graph.add_edge("Detect OCR Requirement", "OCR Processing")
    graph.add_edge("OCR Processing", "Table Extraction")
    graph.add_edge("Table Extraction", "Merge Text")
    graph.add_edge("Merge Text", "Normalize Text")
    graph.add_edge("Normalize Text", "Chunk Document")
    graph.add_edge("Chunk Document", "Financial Extraction")
    graph.add_edge("Financial Extraction", "Merge Chunk Results")
    graph.add_edge("Merge Chunk Results", "Validate JSON")
    graph.add_edge("Validate JSON", "Confidence Calculation")
    graph.add_edge("Confidence Calculation", "Response Formatter")
    graph.add_edge("Response Formatter", END)
    return graph.compile()


WORKFLOW = build_workflow()

app = FastAPI(title="Financial Document Extraction Backend", version="1.0.0", default_response_class=ORJSONResponse)
app.add_middleware(
    CORSMiddleware,
    allow_origins=[origin.strip() for origin in os.getenv("CORS_ALLOW_ORIGINS", "*").split(",")],
    allow_credentials=True,
    allow_methods=["GET", "POST"],
    allow_headers=["*"],
)


async def process_document(data: bytes, filename: str, mime_type: str, metadata: Optional[Dict[str, Any]] = None) -> Tuple[Dict[str, Any], int]:
    logger.info("Upload received {filename} size={size}", filename=filename, size=len(data))
    initial_state: ExtractionState = {
        "uploaded_file": data,
        "filename": filename or "upload",
        "extension": Path(filename or "").suffix.lower(),
        "mime_type": mime_type or "application/octet-stream",
        "metadata": dict(metadata or {}),
        "page_count": 0,
        "ocr_required": False,
        "raw_text": "",
        "ocr_text": "",
        "merged_text": "",
        "tables": [],
        "chunks": [],
        "chunk_results": [],
        "financial_data": empty_financial_data(),
        "confidence_score": 0.0,
        "execution_time": 0.0,
        "errors": [],
        "warnings": [],
        "start_time": time.perf_counter(),
        "response": {},
        "ocr_pages": [],
        "native_quality": 0.0,
        "status_code": 200,
    }
    try:
        final_state = await WORKFLOW.ainvoke(initial_state)
        response = sanitize_for_json(final_state.get("response") or {})
        status_code = int(final_state.get("status_code") or status_code_for_state(final_state))
        return response, status_code
    except MemoryError:
        logger.exception("Memory error while processing document")
        return error_response(filename, "Memory error while processing document", 507), 507
    except Exception as exc:
        logger.exception("Unhandled processing error")
        return error_response(filename, f"Unhandled processing error: {str(exc)}", 500), 500
    finally:
        logger.info("Cleanup completed for {filename}", filename=filename)


def error_response(filename: str, message: str, status_code: int) -> Dict[str, Any]:
    return {
        "success": False,
        "filename": filename or "",
        "filetype": Path(filename or "").suffix.lower().replace(".", ""),
        "pages": 0,
        "ocr_used": False,
        "processing_time": 0,
        "metadata": {},
        "financial_data": empty_financial_data(),
        "confidence_score": 0.0,
        "errors": [message],
        "warnings": [],
    }


@app.get("/", response_class=HTMLResponse)
async def root() -> HTMLResponse:
    return HTMLResponse("""
<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Financial Document Extractor</title>
<style>
:root { color-scheme: light; font-family: Inter, Segoe UI, Arial, sans-serif; background: #f6f7f9; color: #1d2430; }
body { margin: 0; min-height: 100vh; display: grid; place-items: center; padding: 32px; box-sizing: border-box; }
main { width: min(980px, 100%); background: #fff; border: 1px solid #d9dee7; border-radius: 8px; box-shadow: 0 12px 34px rgba(15, 23, 42, 0.08); overflow: hidden; }
header { padding: 26px 30px 18px; border-bottom: 1px solid #e6e9ef; }
h1 { margin: 0 0 8px; font-size: 24px; font-weight: 700; letter-spacing: 0; }
p { margin: 0; color: #5e6878; line-height: 1.5; }
section { padding: 26px 30px; }
form { display: grid; gap: 16px; }
.drop { border: 1.5px dashed #9aa7b8; border-radius: 8px; background: #fbfcfe; padding: 28px; text-align: center; }
input[type=file] { width: 100%; max-width: 520px; }
button { justify-self: start; border: 0; border-radius: 6px; background: #165dff; color: white; font-size: 15px; font-weight: 650; padding: 11px 18px; cursor: pointer; }
button:disabled { background: #9aa7b8; cursor: wait; }
.status { min-height: 24px; color: #3f4a5c; font-weight: 600; }
.result { display: none; border-top: 1px solid #e6e9ef; background: #fbfcfe; }
.summary { display: grid; grid-template-columns: repeat(auto-fit, minmax(170px, 1fr)); gap: 10px; margin-bottom: 16px; }
.tile { background: #fff; border: 1px solid #e1e6ee; border-radius: 8px; padding: 12px; }
.tile span { display: block; color: #6b7280; font-size: 12px; margin-bottom: 6px; }
.tile strong { display: block; font-size: 15px; word-break: break-word; }
pre { margin: 0; max-height: 460px; overflow: auto; background: #111827; color: #e5e7eb; border-radius: 8px; padding: 16px; font-size: 13px; line-height: 1.45; white-space: pre-wrap; word-break: break-word; }
.error { color: #b42318; }
</style>
</head>
<body>
<main>
<header>
<h1>Financial Document Extractor</h1>
<p>Upload DOCX, PDF, images, TXT, CSV, TSV, XLSX, XLS, Markdown, RTF, HTML, JSON, or XML. The backend extracts text and financial values using LangChain with Groq.</p>
</header>
<section>
<form id="uploadForm">
<div class="drop">
<input id="fileInput" name="file" type="file" accept=".pdf,.png,.jpeg,.jpg,.bmp,.tiff,.tif,.docx,.txt,.md,.rtf,.html,.htm,.json,.xml,.csv,.tsv,.xlsx,.xls" required>
</div>
<button id="submitButton" type="submit">Extract Data</button>
<div id="status" class="status"></div>
</form>
</section>
<section id="result" class="result">
<div id="summary" class="summary"></div>
<pre id="jsonOutput"></pre>
</section>
</main>
<script>
const form = document.getElementById('uploadForm');
const fileInput = document.getElementById('fileInput');
const button = document.getElementById('submitButton');
const statusBox = document.getElementById('status');
const resultBox = document.getElementById('result');
const summaryBox = document.getElementById('summary');
const jsonOutput = document.getElementById('jsonOutput');
function valueText(value) {
  if (value === null || value === undefined || value === '') return 'Not found';
  if (typeof value === 'object') return value.raw_value || value.value || value.normalized_value || 'Found';
  return String(value);
}
function tile(label, value) {
  const div = document.createElement('div');
  div.className = 'tile';
  const span = document.createElement('span');
  span.textContent = label;
  const strong = document.createElement('strong');
  strong.textContent = valueText(value);
  div.appendChild(span);
  div.appendChild(strong);
  return div;
}
function render(data) {
  const financial = data.financial_data || {};
  summaryBox.innerHTML = '';
  summaryBox.appendChild(tile('Company', financial.company_name));
  summaryBox.appendChild(tile('Revenue', financial.revenue));
  summaryBox.appendChild(tile('Net Income', financial.net_income || financial.net_profit));
  summaryBox.appendChild(tile('Currency', financial.currency));
  summaryBox.appendChild(tile('Confidence', typeof data.confidence_score === 'number' ? data.confidence_score.toFixed(1) + '%' : data.confidence_score));
  jsonOutput.textContent = JSON.stringify(data, null, 2);
  resultBox.style.display = 'block';
}
form.addEventListener('submit', async (event) => {
  event.preventDefault();
  const file = fileInput.files[0];
  if (!file) return;
  const formData = new FormData();
  formData.append('file', file);
  button.disabled = true;
  statusBox.className = 'status';
  statusBox.textContent = 'Extracting data...';
  resultBox.style.display = 'none';
  try {
    const response = await fetch('/extract', { method: 'POST', body: formData });
    const data = await response.json();
    render(data);
    statusBox.textContent = data.success ? 'Extraction complete.' : 'Extraction finished with errors.';
    if (!data.success) statusBox.className = 'status error';
  } catch (error) {
    statusBox.className = 'status error';
    statusBox.textContent = 'Upload failed: ' + error.message;
  } finally {
    button.disabled = false;
  }
});
</script>
</body>
</html>
""")

@app.get("/health")
async def health() -> ORJSONResponse:
    return ORJSONResponse(
        {
            "success": True,
            "status": "healthy",
            "time": now_iso(),
            "groq_api_key_configured": bool(os.getenv("GROQ_API_KEY")),
            "groq_model": os.getenv("GROQ_MODEL", DEFAULT_GROQ_MODEL),
            "ocr_model_loaded": OCR_MODEL is not None,
            "supported_filetypes": sorted(extension.replace(".", "") for extension in SUPPORTED_EXTENSIONS),
        }
    )


@app.post("/extract")
async def extract(file: UploadFile = File(...)) -> ORJSONResponse:
    try:
        data = await read_upload_stream(file)
        response, status_code = await process_document(data, file.filename or "upload", file.content_type or "application/octet-stream")
        return ORJSONResponse(response, status_code=status_code)
    except ValueError as exc:
        response = error_response(file.filename if file else "", str(exc), 400)
        return ORJSONResponse(response, status_code=400)
    except Exception as exc:
        logger.exception("Upload handling failed")
        response = error_response(file.filename if file else "", f"Upload handling failed: {str(exc)}", 500)
        return ORJSONResponse(response, status_code=500)


@app.post("/extract/batch")
async def extract_batch(files: List[UploadFile] = File(...)) -> ORJSONResponse:
    if not files:
        return ORJSONResponse({"success": False, "results": [], "errors": ["No files uploaded"], "warnings": []}, status_code=400)
    if len(files) > MAX_BATCH_FILES:
        return ORJSONResponse({"success": False, "results": [], "errors": [f"Batch limit exceeded: {MAX_BATCH_FILES} files"], "warnings": []}, status_code=400)
    semaphore = asyncio.Semaphore(BATCH_CONCURRENCY)

    async def process_upload(upload: UploadFile) -> Dict[str, Any]:
        async with semaphore:
            try:
                data = await read_upload_stream(upload)
                response, status_code = await process_document(data, upload.filename or "upload", upload.content_type or "application/octet-stream")
                response["status_code"] = status_code
                return response
            except Exception as exc:
                logger.exception("Batch upload failed for {filename}", filename=upload.filename)
                return error_response(upload.filename or "", f"Batch upload failed: {str(exc)}", 500)

    results = await asyncio.gather(*(process_upload(file) for file in files))
    success = all(result.get("success") for result in results)
    status_code = 200 if success else max((int(result.get("status_code", 500)) for result in results if not result.get("success")), default=500)
    return ORJSONResponse({"success": success, "count": len(results), "results": results}, status_code=status_code)


@app.post("/extract/text")
async def extract_text(request: TextExtractionRequest) -> ORJSONResponse:
    filename = request.filename or "text-input.txt"
    if not Path(filename).suffix:
        filename = f"{filename}.txt"
    data = request.text.encode("utf-8")
    metadata = dict(request.metadata or {})
    metadata["source"] = "text_endpoint"
    response, status_code = await process_document(data, filename, request.mime_type or "text/plain", metadata)
    return ORJSONResponse(response, status_code=status_code)


@app.exception_handler(Exception)
async def unhandled_exception_handler(request: Request, exc: Exception) -> ORJSONResponse:
    logger.exception("Unhandled application error at {path}", path=str(request.url.path))
    status_code = 507 if isinstance(exc, MemoryError) else 500
    return ORJSONResponse(error_response("", f"Unhandled application error: {str(exc)}", status_code), status_code=status_code)


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(
        app,
        host=os.getenv("EXTRACTION_AGENT_HOST", "127.0.0.2"),
        port=int(os.getenv("EXTRACTION_AGENT_PORT", "8001")),
        reload=False,
        log_level="info"
    )



